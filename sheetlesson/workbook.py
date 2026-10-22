"""Typed student workbook export and deliberately limited XLSX migration."""

# Standard Library
import json
import math
import re
import pathlib
import zipfile

# PIP modules
import defusedxml.ElementTree
import openpyxl
import openpyxl.styles
import openpyxl.workbook.properties

# Local modules
import sheetlesson.model


MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
MAX_IMPORT_CELLS = 250000


#============================================
def _xml_cells(path: pathlib.Path) -> dict[str, dict[str, tuple[str, bool]]]:
	"""Read storage types that openpyxl collapses, notably explicit empty strings."""
	result = {}
	with zipfile.ZipFile(path) as archive:
		for name in archive.namelist():
			if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
				continue
			root = defusedxml.ElementTree.fromstring(archive.read(name))
			cells = {}
			for cell in root.iter(f"{MAIN_NS}c"):
				kind = cell.attrib.get("t", "n")
				cells[cell.attrib["r"]] = (kind, cell.find(f"{MAIN_NS}f") is not None)
			result[name] = cells
	return result


#============================================
def _sheet_parts(path: pathlib.Path) -> list[str]:
	"""Resolve worksheet parts in workbook order rather than guessing sheet IDs."""
	with zipfile.ZipFile(path) as archive:
		root = defusedxml.ElementTree.fromstring(archive.read("xl/workbook.xml"))
		rels = defusedxml.ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
		mapping = {item.attrib["Id"]: item.attrib["Target"] for item in rels}
		parts = []
		for sheet in root.find(f"{MAIN_NS}sheets"):
			rel_id = sheet.attrib[
				"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
			target = mapping[rel_id]
			part = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
			parts.append(part)
	return parts


#============================================
def _logical_cells(path: pathlib.Path) -> tuple[list[str], list[dict]]:
	"""Read logical values without formula caches, retaining text and blank types."""
	storage = _xml_cells(path)
	parts = _sheet_parts(path)
	book = openpyxl.load_workbook(path, data_only=False)
	all_cells = []
	for sheet, part in zip(book.worksheets, parts, strict=True):
		cells = {}
		for address, (kind, formula) in storage[part].items():
			cell = sheet[address]
			value = cell.value
			if formula:
				scalar = sheetlesson.model.Scalar("formula", value)
			elif kind in ("s", "inlineStr", "str"):
				scalar = sheetlesson.model.Scalar("text", "" if value is None else value)
			elif value is None:
				scalar = sheetlesson.model.Scalar("blank", None)
			elif isinstance(value, bool):
				scalar = sheetlesson.model.Scalar("boolean", value)
			elif isinstance(value, (int, float)):
				if not math.isfinite(value):
					raise ValueError(f"{sheet.title}!{address}: nonfinite number is unsupported")
				scalar = sheetlesson.model.Scalar("number", value)
			else:
				raise ValueError(f"{sheet.title}!{address}: unsupported cell value {value!r}")
			cells[(cell.row, cell.column)] = scalar
		all_cells.append(cells)
	book.close()
	return book.sheetnames, all_cells


#============================================
def write_workbook(document: sheetlesson.model.Document, output_path: pathlib.Path | str) -> None:
	"""Write only authored seeds, with distinct styling for empty student targets."""
	if document.has_errors:
		raise ValueError("Cannot export a document with compilation errors")
	if not document.worksheets:
		raise ValueError("A workbook requires at least one worksheet")
	book = openpyxl.Workbook()
	book.remove(book.active)
	book.calculation = openpyxl.workbook.properties.CalcProperties(
		calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode="auto")
	for source in document.worksheets:
		sheet = book.create_sheet(source.name)
		widths = {}
		for (row, column), seed in source.cells.items():
			cell = sheet.cell(row, column)
			scalar = seed.scalar
			if scalar.kind == "number" and not math.isfinite(scalar.value):
				raise ValueError(f"{source.name}!{cell.coordinate}: nonfinite number")
			if scalar.kind == "text" and len(scalar.value) > 32767:
				raise ValueError(f"{source.name}!{cell.coordinate}: text exceeds XLSX 32767 characters")
			cell.value = scalar.value
			if scalar.kind == "text":
				# Leading equals signs and explicit empty strings remain literal text.
				cell.data_type = "s"
			cell.alignment = openpyxl.styles.Alignment(vertical="top", wrap_text=True)
			length = max((len(line) for line in str(scalar.value).splitlines()), default=0)
			widths[column] = min(48, max(widths.get(column, 12), length + 2))
		for region in source.regions:
			for column in range(region.start_column, region.end_column + 1):
				sheet.cell(region.start_row, column).font = openpyxl.styles.Font(bold=True)
		for row, column in source.entry_cells:
			seed = source.cells.get((row, column))
			if seed is None or seed.scalar.kind == "blank":
				cell = sheet.cell(row, column)
				cell.fill = openpyxl.styles.PatternFill("solid", fgColor="FFF2CC")
				cell.border = openpyxl.styles.Border(
					bottom=openpyxl.styles.Side(style="thin", color="9C6500"))
		for column, width in widths.items():
			sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
	book.save(output_path)
	book.close()
	_preserve_numbers(document, pathlib.Path(output_path))


#============================================
def _preserve_numbers(document: sheetlesson.model.Document, path: pathlib.Path) -> None:
	"""Avoid openpyxl's 16-digit formatting changing a finite Python numeric value."""
	parts = _sheet_parts(path)
	with zipfile.ZipFile(path) as archive:
		members = [(item, archive.read(item.filename)) for item in archive.infolist()]
	updated = {}
	for sheet, part in zip(document.worksheets, parts, strict=True):
		root = defusedxml.ElementTree.fromstring(dict(
			(item.filename, content) for item, content in members)[part])
		for cell in root.iter(f"{MAIN_NS}c"):
			row, column = openpyxl.utils.cell.coordinate_to_tuple(cell.attrib["r"])
			seed = sheet.cells.get((row, column))
			if seed is not None and seed.scalar.kind == "number":
				cell.find(f"{MAIN_NS}v").text = repr(seed.scalar.value)
		updated[part] = defusedxml.ElementTree.tostring(root, encoding="utf-8")
	with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
		for item, content in members:
			archive.writestr(item, updated.get(item.filename, content))


#============================================
def verify_workbook(document: sheetlesson.model.Document, path: pathlib.Path | str) -> None:
	"""Raise on changed worksheet order, seed values, types, or unexpected content."""
	names, observed = _logical_cells(pathlib.Path(path))
	if names != [sheet.name for sheet in document.worksheets]:
		raise ValueError("Workbook verification failed: worksheet order or names changed")
	blank = sheetlesson.model.Scalar("blank", None)
	for sheet, cells in zip(document.worksheets, observed, strict=True):
		for address, seed in sheet.cells.items():
			if cells.get(address, blank) != seed.scalar:
				raise ValueError(f"Workbook verification failed: {sheet.name}!{address} changed")
		for address, scalar in cells.items():
			if address not in sheet.cells and scalar.kind != "blank":
				raise ValueError(f"Workbook verification failed: unexpected {sheet.name}!{address}")


#============================================
def _inventory(path: pathlib.Path) -> None:
	"""Block content whose meaning cannot survive the supported Markdown format."""
	issues = []
	blocked_tags = {
		"mergeCells", "legacyDrawing", "legacyDrawingHF", "oleObjects",
		"controls", "tableParts", "pivotTableParts", "dataValidations", "conditionalFormatting",
		"hyperlinks", "extLst", "sheetProtection", "autoFilter", "scenarios",
		"definedName",
	}
	with zipfile.ZipFile(path) as archive:
		for name in archive.namelist():
			if any(piece in name for piece in (
				"vbaProject", "xl/charts/", "xl/pivot", "xl/externalLinks/", "xl/comments",
				"xl/threadedComments/", "xl/embeddings/", "xl/media/")):
				issues.append(f"unsupported package content: {name}")
			if name.startswith("xl/drawings/") and name.endswith(".xml"):
				# Excel emits empty drawing containers; these contain no lost content.
				if len(defusedxml.ElementTree.fromstring(archive.read(name))):
					issues.append(f"unsupported drawing content: {name}")
			if name == "xl/workbook.xml" or name.startswith("xl/worksheets/sheet"):
				if not name.endswith(".xml"):
					continue
				root = defusedxml.ElementTree.fromstring(archive.read(name))
				for item in root.iter():
					tag = item.tag.rsplit("}", 1)[-1]
					if tag in blocked_tags:
						issues.append(f"{name}: unsupported {tag}")
					if tag == "workbookProtection" and item.attrib:
						issues.append(f"{name}: workbook protection")
					if (tag == "calcPr"
						and item.attrib.get("iterate", "0").lower() in ("1", "true")):
						issues.append(
							"workbook enables iterative calculation; formula behavior cannot be preserved safely")
					if tag == "sheet" and item.attrib.get("state", "visible") != "visible":
						issues.append(f"hidden worksheet: {item.attrib['name']}")
					if tag in ("row", "col") and item.attrib.get("hidden") in ("1", "true"):
						issues.append(f"{name}: hidden {tag} content")
					if tag == "f" and item.attrib.get("t", "normal") not in ("normal", "shared"):
						issues.append(f"{name}: unsupported formula type {item.attrib['t']}")
					if tag == "c" and item.attrib.get("t") in ("e", "d"):
						issues.append(f"{name}!{item.attrib['r']}: error or date cell")
			if name == "xl/sharedStrings.xml" or name.startswith("xl/worksheets/sheet"):
				if name.endswith(".xml"):
					root = defusedxml.ElementTree.fromstring(archive.read(name))
					if any(True for item in root.iter(f"{MAIN_NS}r")):
						issues.append(f"{name}: rich text")
	if issues:
		message = "XLSX import blocked; supply a simpler workbook or extend the format:\n"
		message += "\n".join(sorted(set(issues)))
		raise ValueError(message)
	# Date styles may encode numeric cells as dates, even without the ISO date type.
	book = openpyxl.load_workbook(path, data_only=False)
	if book.epoch == openpyxl.utils.datetime.CALENDAR_MAC_1904:
		issues.append("workbook uses the 1904 date epoch; numeric serials cannot be preserved safely")
	storage = _xml_cells(path)
	for sheet, part in zip(book.worksheets, _sheet_parts(path), strict=True):
		for address in storage[part]:
			if sheet[address].is_date:
				issues.append(f"{sheet.title}!{address}: date/time cell")
		if sheet.title.startswith("'") or sheet.title.endswith("'"):
			issues.append(f"worksheet name cannot be represented: {sheet.title!r}")
		if (sheet.title != sheet.title.strip() or "\n" in sheet.title or "\r" in sheet.title
			or re.search(r"\s#+\s*$", sheet.title)):
			issues.append(f"worksheet name needs normalization before import: {sheet.title!r}")
	book.close()
	if issues:
		raise ValueError("XLSX import blocked; simplify these cells:\n" + "\n".join(issues))


#============================================
def _markdown_scalar(scalar: sheetlesson.model.Scalar) -> str:
	"""Encode a scalar inside a pipe table without changing its semantic type."""
	if scalar.kind == "blank":
		return ""
	if scalar.kind == "boolean":
		return "TRUE" if scalar.value else "FALSE"
	if scalar.kind == "number":
		return str(scalar.value)
	value = json.dumps(scalar.value, ensure_ascii=True) if scalar.kind == "text" else scalar.value
	# Code spans preserve literal Markdown, entities and backslashes. GFM pipes still escape.
	if scalar.kind == "formula" and ("\n" in value or "\r" in value):
		raise ValueError("Multiline formulas require a simpler single-line source before import")
	runs = [len(piece) for piece in re.findall(r"`+", value)]
	fence = "`" * (max(runs, default=0) + 1)
	value = value.replace("|", "\\|")
	result = f"{fence} {value} {fence}"
	return result


#============================================
def _row_runs(cells: dict) -> list[tuple[int, int, int, int]]:
	"""Bound occupied consecutive row runs without allocating empty worksheet extents."""
	occupied = {key: value for key, value in cells.items() if value.kind != "blank"}
	if not occupied:
		return [(1, 1, 1, 1)]
	row_bounds = {}
	for row, column in occupied:
		low, high = row_bounds.get(row, (column, column))
		row_bounds[row] = (min(low, column), max(high, column))
	regions = []
	for row in sorted(row_bounds):
		low, high = row_bounds[row]
		if regions and regions[-1][2] + 1 == row:
			start, left, end, right = regions[-1]
			regions[-1] = (start, min(left, low), row, max(right, high))
		else:
			regions.append((row, low, row, high))
	if len(regions) == 1:
		start_row, start_column, end_row, end_column = regions[0]
		# Keep small near-A1 tables intuitive while retaining all original coordinates.
		if start_row <= 10 and start_column <= 10:
			start_row, start_column = 1, 1
		regions[0] = (start_row, start_column, end_row, end_column)
	area = sum((end_row - row + 1) * (end_column - column + 1)
		for row, column, end_row, end_column in regions)
	if area > MAX_IMPORT_CELLS:
		raise ValueError(
			f"XLSX sparse extents require {area} table cells; limit is {MAX_IMPORT_CELLS}. "
			"Split widely separated columns into simpler worksheets before importing.")
	return regions


#============================================
def import_workbook(path: pathlib.Path | str) -> str:
	"""Create a draft single-source document only after a loss-prevention inventory."""
	path = pathlib.Path(path)
	_inventory(path)
	names, sheets = _logical_cells(path)
	title = path.stem.replace("_", " ").replace("\n", " ")
	lines = [f"# {title}", "", "> Import summary: Draft migration; instructions require authoring.",
		"> Ordinary fonts, fills, widths, and number formats are discarded.",
		"> Formulas are preserved; calculation was not performed and cached results are not imported.",
		"", "## Workbook", ""]
	blank = sheetlesson.model.Scalar("blank", None)
	for name, cells in zip(names, sheets, strict=True):
		lines.extend([f"### Sheet: {name}", ""])
		for index, (row, column, end_row, end_column) in enumerate(_row_runs(cells)):
			if index > 0 or (row, column) != (1, 1):
				address = f"{openpyxl.utils.get_column_letter(column)}{row}"
				lines.extend([f"#### At: {address}", ""])
			for source_row in range(row, end_row + 1):
				values = [_markdown_scalar(cells.get((source_row, col), blank))
					for col in range(column, end_column + 1)]
				lines.append("| " + " | ".join(values) + " |")
				if source_row == row:
					lines.append("| " + " | ".join("---" for value in values) + " |")
			lines.append("")
	lines.extend(["## Tutorial", "", "Write the tutorial instructions here.", ""])
	result = "\n".join(lines)
	return result
