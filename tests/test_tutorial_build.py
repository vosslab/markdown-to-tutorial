"""Protect published teaching bundles from failed builds and user-edit loss.

Failures mean publication protection regressed: fix staging/ownership before releasing.
Conversion is replaced locally so these integrity checks stay offline and fast.
"""

# Standard library
import json
import io
import pathlib
import subprocess

# Third-party modules
import docx
import pytest
import pypdf

# Local modules
import sheetlesson.build
import sheetlesson.model
import sheetlesson.pdf_export


#============================================
def make_document(directory: pathlib.Path) -> sheetlesson.model.Document:
	"""Create a minimal compiled model for publication integrity tests."""
	source = directory / "lesson.md"
	source.write_text("A tutorial\n", encoding="utf-8")
	document = sheetlesson.model.Document(source, source.read_text(encoding="utf-8"))
	span = sheetlesson.model.SourceSpan(str(source), 1, 1)
	sheet = sheetlesson.model.Worksheet("Data", span)
	sheet.cells[(1, 1)] = sheetlesson.model.CellSeed(
		1, 1, sheetlesson.model.Scalar("text", "Value"), span
	)
	document.worksheets.append(sheet)
	return document


#============================================
def write_docx(document: sheetlesson.model.Document, path: pathlib.Path) -> None:
	"""Provide a valid local package without exercising the independent renderer."""
	output = docx.Document()
	output.add_paragraph(document.source)
	output.save(path)


#============================================
def write_pdf(source: pathlib.Path, destination: pathlib.Path) -> None:
	"""Stand in for conversion without invoking LibreOffice in pytest."""
	if not source.is_file():
		raise FileNotFoundError(source)
	destination.write_bytes(b"%PDF-1.4\nMock conversion evidence only\n")


#============================================
def fail_pdf(source: pathlib.Path, destination: pathlib.Path) -> None:
	"""Represent a converter failure after other staged artifacts exist."""
	raise RuntimeError("conversion failed")


#============================================
def test_failed_conversion_preserves_bundle_and_subset_replaces_it(
	tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
	document = make_document(tmp_path)
	monkeypatch.setattr(sheetlesson.build.sheetlesson.docx_renderer, "write_document", write_docx)
	monkeypatch.setattr(sheetlesson.build.sheetlesson.pdf_export, "export_pdf", write_pdf)
	destination = tmp_path / "published"
	sheetlesson.build.build_document(document, destination)
	before = {path.name: path.read_bytes() for path in destination.iterdir()}
	monkeypatch.setattr(sheetlesson.build.sheetlesson.pdf_export, "export_pdf", fail_pdf)
	with pytest.raises(RuntimeError, match="conversion failed"):
		sheetlesson.build.build_document(document, destination)
	assert {path.name: path.read_bytes() for path in destination.iterdir()} == before
	sheetlesson.build.build_document(document, destination, "xlsx")
	assert set(path.suffix for path in destination.iterdir()) == {".xlsx", ".json"}
	report = json.loads((destination / "validation.json").read_text(encoding="utf-8"))
	assert report["external_behavior"]["google_import"] == "unchecked"


#============================================
def test_modified_artifact_and_foreign_directory_are_preserved(tmp_path: pathlib.Path) -> None:
	document = make_document(tmp_path)
	destination = tmp_path / "published"
	sheetlesson.build.build_document(document, destination, "xlsx")
	workbook = destination / "lesson.xlsx"
	workbook.write_bytes(b"user edits")
	with pytest.raises(ValueError, match="modified"):
		sheetlesson.build.build_document(document, destination, "xlsx")
	assert workbook.read_bytes() == b"user edits"
	foreign = tmp_path / "foreign"
	foreign.mkdir()
	(foreign / "notes.txt").write_text("keep", encoding="utf-8")
	with pytest.raises(ValueError, match="not an owned"):
		sheetlesson.build.build_document(document, foreign, "xlsx")
	assert (foreign / "notes.txt").read_text(encoding="utf-8") == "keep"


#============================================
def test_successful_converter_with_unreadable_pdf_preserves_bundle(
	tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
	"""A converter's zero exit and PDF signature must not publish a corrupt bundle."""
	document = make_document(tmp_path)
	monkeypatch.setattr(sheetlesson.build.sheetlesson.docx_renderer, "write_document", write_docx)
	monkeypatch.setattr(sheetlesson.pdf_export.shutil, "which", lambda name: "/mock/soffice")
	writer = pypdf.PdfWriter()
	writer.add_blank_page(width=100, height=100)
	buffer = io.BytesIO()
	writer.write(buffer)
	pdf_bytes = buffer.getvalue()

	def convert(command: list[str], **options: object) -> subprocess.CompletedProcess[str]:
		output = pathlib.Path(command[-2]) / pathlib.Path(command[-1]).with_suffix(".pdf").name
		output.write_bytes(pdf_bytes)
		return subprocess.CompletedProcess(command, 0, stdout="converted", stderr="")

	monkeypatch.setattr(sheetlesson.pdf_export.subprocess, "run", convert)
	destination = tmp_path / "published"
	sheetlesson.build.build_document(document, destination)
	before = {path.name: path.read_bytes() for path in destination.iterdir()}
	pdf_bytes = b"%PDF-"
	with pytest.raises(RuntimeError, match="unreadable PDF"):
		sheetlesson.build.build_document(document, destination)
	assert {path.name: path.read_bytes() for path in destination.iterdir()} == before
