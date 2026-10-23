"""Preserve ordinary instructional prose with explicit executable link bindings."""

import json
import docx
import sheets_validation_lib
import tutorial_lib.compiler
import tutorial_lib.docx_styles
import tutorial_lib.docx_renderer


def compile_bindings(instructions: str, metadata: str = "") -> object:
	source = "# Original lesson\n\n## Workbook\n\n### Sheet: Data\n\n"
	source += metadata + "\n| Sleep | Score |\n| --- | --- |\n| 6 | 90 |\n| 7 | 95 |\n\n"
	source += "## Tutorial\n\n### **PART: Quadratic Regression**\n\n"
	source += "#### **Step: Create a SQUARED Sleep-Hours Column**\n\n"
	return tutorial_lib.compiler.compile_source(source + instructions)


def test_numbered_prose_bindings_replay_every_operation_in_order() -> None:
	"""A compact four-item explanation must execute without rewriting its prose."""
	instructions = "To let our regression model curve, we need to square our predictor variable (x).\n\n"
	instructions += "1. Add a column beside your data.\n"
	instructions += "2. Label it [Sleep squared](label:C1).\n"
	instructions += "3. Enter [=A2^2](enter:C2), then [fill down](fill:C2:C3).\n"
	instructions += "4. Check [the squared values](select:C2:C3). Don't worry, they didn't sleep for 36 hours!\n"
	document = compile_bindings(instructions)
	assert not document.has_errors
	assert document.source.endswith(instructions)
	assert [action.kind for action in document.actions] == ["label", "enter", "fill", "select"]
	assert document.headings[0].title == "Quadratic Regression"
	assert document.headings[1].title == "Create a SQUARED Sleep-Hours Column"
	assert document.headings[1].prefix_length == len("Step: ")
	assert all("action" not in token.meta for token in document.tokens)
	legacy = compile_bindings('- **Label C1:** Sleep squared\n- **Enter C2:** `=A2^2`\n'
		'- **Fill C2:C3:** Down\n- **Select:** C2:C3\n')
	bound_writes = sheets_validation_lib.prepare_cases(document)[0]["writes"]
	legacy_writes = sheets_validation_lib.prepare_cases(legacy)[0]["writes"]
	assert [(write["address"], write["value"]) for write in bound_writes] == [
		(write["address"], write["value"]) for write in legacy_writes]
	# Both assignments in one inline must be applied before the following replacement.
	ordered = compile_bindings('Put [1](enter:C2), [2](enter:D2), then [3](replace:C2).')
	assert not ordered.has_errors
	assert [write["address"] for write in sheets_validation_lib.prepare_cases(ordered)[0]["writes"]] == [
		"C2", "D2", "C2"]


def test_literal_binding_and_formula_source_remain_distinct() -> None:
	"""Literal copy text cannot accidentally become an unknown-sheet input or formula."""
	for label in ('`Absent!A1`', '=Absent!A1*3*4'):
		document = compile_bindings(f'Use [{label}](label:C2).')
		assert not document.has_errors
		assert document.actions[0].scalar.kind == "text"
		assert not document.formula_checks
	for value in ('=Absent!A1*3*4', '*literal* _text_ <br> a\\b', 'say "hello"'):
		document = compile_bindings(f'Enter [{json.dumps(value)}](enter:C2).')
		assert not document.has_errors
		assert document.actions[0].scalar.value == value
		assert not document.formula_checks
	formula = '=IF(A2>0,"A*B",A2*3*4)'
	document = compile_bindings(f'Enter [{formula}](enter:C2).')
	assert not document.has_errors
	assert document.actions[0].scalar.value == formula
	assert document.formula_checks[0].formula == formula
	qualified = compile_bindings('Enter [4](enter:%27Data%27%21C2).')
	assert not qualified.has_errors
	assert qualified.actions[0].target.sheet == "Data"


def test_bindings_and_explicit_role_regions_fail_closed() -> None:
	"""Broken machine targets and conflicting role metadata must stop a build."""
	for binding in ('[1](enter:A0)', '[1](label:A1:B2)', '[down](fill:C2)',
		'[down](fill:C3:C2)', '[1](enter:Absent!C2)', '["bad](enter:C2)'):
		assert compile_bindings(binding).has_errors
	document = compile_bindings('Read [the source](https://example.org).',
		'- **Region A1:B3:** source\n- **Region C1:D3:** calculation\n')
	assert not document.has_errors
	assert [region.role for region in document.worksheets[0].role_regions] == ["source", "calculation"]
	for metadata in ('- **Region A0:B3:** source\n', '- **Region C1:C3:** answer\n',
		'- **Region A1:B3:** source\n- **Region B2:C3:** results\n'):
		assert compile_bindings('Read the data.', metadata).has_errors


def test_binding_emphasis_preserves_visible_prose_and_scalar_values() -> None:
	"""Author emphasis must survive DOCX without becoming workbook cell characters."""
	formula = '=IF(A2>0,"A_B*C",A2*3*4)'
	literal = 'literal *stars* _under_score_ \\path'
	instructions = 'Label [**Sample *Size***](label:D5).\n\n'
	instructions += 'Enter [*3.5*](enter:D6), then [**' + formula + '**](enter:D7).\n\n'
	instructions += 'Keep [' + json.dumps(literal) + '](enter:D8).\n\n'
	instructions += r'Label [\*literal\* and **bold**](label:D9).' + '\n\n'
	instructions += 'Keep [**' + json.dumps(literal) + '**](enter:D10).'
	document = compile_bindings(instructions)
	assert not document.has_errors
	assert document.source.endswith(instructions)
	assert [action.scalar.value for action in document.actions] == [
		'Sample Size', 3.5, formula, literal, '*literal* and bold', literal]
	assert document.formula_checks[0].formula == formula
	writes = sheets_validation_lib.prepare_cases(document)[0]['writes']
	assert [write['value'] for write in writes] == [
		{'stringValue': 'Sample Size'}, {'numberValue': 3.5}, {'formulaValue': formula},
		{'stringValue': literal}, {'stringValue': '*literal* and bold'}, {'stringValue': literal}]
	output = docx.Document()
	tutorial_lib.docx_styles.configure(output)
	tutorial_lib.docx_renderer.render_tokens(output, document)
	paragraphs = output.paragraphs[-5:]
	assert [paragraph.text for paragraph in paragraphs] == [
		'Label Sample Size.', f'Enter 3.5, then {formula}.',
		'Keep ' + json.dumps(literal) + '.', 'Label *literal* and bold.',
		'Keep ' + json.dumps(literal) + '.']
	assert any(run.text == 'Size' and run.bold and run.italic for run in paragraphs[0].runs)
	assert any(run.text == '3.5' and run.italic for run in paragraphs[1].runs)
	assert any(run.text == formula and run.bold for run in paragraphs[1].runs)
	assert any(run.text == json.dumps(literal) and run.bold for run in paragraphs[-1].runs)


def test_quoted_label_markdown_and_json_scalars_have_distinct_contexts() -> None:
	"""Quoted prose labels render Markdown; JSON assignments retain literal copy text."""
	instructions = 'Label ["Significance Level (&alpha;)."](label:D14).\n\n'
	instructions += 'Label ["*Two-Tail Decision*."](label:D16).\n\n'
	instructions += 'Label [`&alpha; *literal*`](label:D17).\n\n'
	instructions += 'Label ["=Absent!A1*3*4"](label:D18).\n\n'
	instructions += 'Label [=Absent!A1*3*4](label:D19).\n\n'
	literal = '&alpha; *literal* _text_ \\path'
	instructions += 'Enter [' + json.dumps(literal) + '](enter:D20).\n\n'
	instructions += 'Replace [**' + json.dumps(literal) + '**](replace:D20).'
	document = compile_bindings(instructions)
	assert not document.has_errors
	values = ['"Significance Level (\u03b1)."', '"Two-Tail Decision."',
		'&alpha; *literal*', '"=Absent!A1*3*4"', '=Absent!A1*3*4', literal, literal]
	assert [action.scalar.value for action in document.actions] == values
	assert not document.formula_checks
	writes = sheets_validation_lib.prepare_cases(document)[0]['writes']
	assert [write['value'] for write in writes] == [{'stringValue': value} for value in values]
	output = docx.Document()
	tutorial_lib.docx_styles.configure(output)
	tutorial_lib.docx_renderer.render_tokens(output, document)
	paragraphs = output.paragraphs[-7:]
	assert [paragraph.text for paragraph in paragraphs] == [
		'Label ' + value + '.' for value in values[:5]] + [
		'Enter ' + json.dumps(literal) + '.', 'Replace ' + json.dumps(literal) + '.']
	assert any(run.text == 'Two-Tail Decision' and run.italic for run in paragraphs[1].runs)
	assert any(run.text == json.dumps(literal) and run.bold for run in paragraphs[-1].runs)
	for token in document.tokens:
		assert not any(child.meta.get('role') == 'formula' for child in token.children or [])


def test_narrative_reference_to_same_paragraph_destination_is_not_an_input_read() -> None:
	"""Naming a cell being assigned in prose must not imply it was read first."""
	instructions = 'In cell `I2`, enter [`=COUNT(A2:B2)`](enter:I2).'
	document = compile_bindings(instructions)
	assert not any(diagnostic.category == "unavailable-input" for diagnostic in document.diagnostics)
	assert document.formula_checks[0].status == "verified"

	missing = compile_bindings('In cell `I2`, enter [`=A99`](enter:I2).')
	assert any(diagnostic.category == "unavailable-input"
		and diagnostic.address == "A99" for diagnostic in missing.diagnostics)


def test_formula_inputs_keep_order_across_multiple_bindings_in_one_paragraph() -> None:
	"""Destination prose is exempt, while a formula still sees earlier writes only."""
	future_input = compile_bindings(
		'In cell `I2`, enter [`=J2`](enter:I2), then put [4](enter:J2) in cell `J2`.')
	assert any(diagnostic.category == "unavailable-input"
		and diagnostic.address == "J2" for diagnostic in future_input.diagnostics)

	prior_input = compile_bindings(
		'Put [4](enter:J2) in cell `J2`, then enter [`=J2`](enter:I2) in cell `I2`.')
	assert not any(diagnostic.category == "unavailable-input"
		for diagnostic in prior_input.diagnostics)
