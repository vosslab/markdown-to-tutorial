"""Protect native formula fidelity and document navigation in DOCX output."""

import io
import pathlib
import base64
import zipfile
import copy
import json

import markdown_it
import pytest
import PIL.Image
import PIL.features
import defusedxml.ElementTree
import docx

import tutorial_lib.model
import tutorial_lib.parser
import tutorial_lib.compiler
import tutorial_lib.docx_renderer
import tutorial_lib.docx_styles


NAMESPACES = {
	"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
	"r": "http://schemas.openxmlformats.org/package/2006/relationships",
}


#============================================
def test_managed_screenshot_comments_are_nonvisible_metadata_only(tmp_path: pathlib.Path) -> None:
	"""Accept screenshot boundaries without allowing HTML or hiding literal examples."""
	begin = '<!-- screenshots:begin (managed by screenshot-docs) -->'
	end = '<!-- screenshots:end -->'
	prefix = '# Metadata\n\n## Workbook\n\n### Sheet: Data\n\n'
	prefix += '| Value |\n| --- |\n| 1 |\n\n## Tutorial\n\n### Part: Review\n\n'
	prefix += '#### Step: Observe\n\n'
	source = prefix + '  ' + begin + '  \n' + end + '\n\nRead the evidence.\n'
	document = tutorial_lib.compiler.compile_source(source)
	assert not document.has_errors
	assert document.source == source
	path = tmp_path / 'metadata.docx'
	tutorial_lib.docx_renderer.write_document(document, path)
	output = docx.Document(path)
	text = '\n'.join(paragraph.text for paragraph in output.paragraphs)
	assert 'Read the evidence.' in text
	assert 'screenshots:' not in text

	for content in ('`' + begin + '`', '```html\n' + begin + '\n' + end + '\n```'):
		document = tutorial_lib.compiler.compile_source(prefix + content)
		assert not document.has_errors
		output = docx.Document()
		tutorial_lib.docx_styles.configure(output)
		tutorial_lib.docx_renderer.render_tokens(output, document)
		assert begin in '\n'.join(paragraph.text for paragraph in output.paragraphs)

	for content in ('<div>Hidden</div>', '<script>alert(1)</script>', '<!-- ordinary -->',
		begin + '<div>Hidden</div>', begin + '\n<script>alert(1)</script>',
		begin.replace('screenshots:begin', 'screenshots:BEGIN'), 'Text ' + end):
		document = tutorial_lib.compiler.compile_source(prefix + content)
		assert any(item.category == 'unsupported-content' for item in document.diagnostics)


#============================================
def test_authored_backslash_breaks_keep_formula_and_literal_bytes() -> None:
	"""Paragraph and list boundaries consume break markers without corrupting escapes."""
	formula = '=IF(A2="a\\b",1,0)'
	source = '**Strong**\\\ncontinued with ' + formula + '\\\nnext.\n\n'
	source += '- First\\\n- **Second**\\\n- Use ' + formula + '\\\n\n'
	source += '`literal\\` and "JSON a\\\\b".\n\nEscaped slash\\\\\n'
	document = tutorial_lib.parser.parse_source(source, 'breaks.md')
	document.content_token_indices = list(range(len(document.tokens)))
	output = docx.Document()
	tutorial_lib.docx_styles.configure(output)
	tutorial_lib.docx_renderer.render_tokens(output, document)
	assert document.source == source
	assert [paragraph.text for paragraph in output.paragraphs] == [
		'Strong\ncontinued with ' + formula + '\nnext.',
		'First', 'Second', 'Use ' + formula,
		'literal\\ and "JSON a\\b".', 'Escaped slash\\',
	]
	formula_runs = [run for paragraph in output.paragraphs for run in paragraph.runs
		if run.text == formula]
	assert len(formula_runs) == 2
	assert all(run.font.color.rgb == docx.shared.RGBColor.from_string('007A00')
		for run in formula_runs)
	assert len(output._element.findall('.//w:br', NAMESPACES)) == 2


#============================================
def test_automatic_formula_color_and_compact_list_paragraphs() -> None:
	"""All formula routes share copyable green type; lists leave body spacing intact."""
	formula = '=IF(A2<3,"A*B",SUM(A2:B2))'
	source = '# Automatic\n\n## Workbook\n\n### Sheet: Data\n\n'
	source += '| A | B |\n| --- | --- |\n| 1 | 2 |\n\n'
	source += '## Tutorial\n\n### Part: Analysis\n\n#### Step: Calculate\n\n'
	source += 'Body use ' + formula + ', then compare.\n\n'
	source += 'Body quote "' + formula + '".\n\n'
	source += '- **Enter C2:** ' + formula + '\n\n'
	source += '  Continuation with `' + formula + '`.\n\n'
	source += '  - **Explain:** [' + formula + '](https://example.org)\n\n'
	source += '- **Enter C3:**\n\n  ```sheets\n  ' + formula + '\n  ```\n\n'
	source += '| Note | Formula |\n| --- | --- |\n| Compare | ' + formula + ' |\n'
	source += '| Quoted | "' + formula + '" |\n'
	document = tutorial_lib.compiler.compile_source(source)
	assert not document.has_errors
	before = copy.deepcopy([token.as_dict() for token in document.tokens])
	output = docx.Document()
	tutorial_lib.docx_styles.configure(output)
	tutorial_lib.docx_renderer.render_tokens(output, document)
	assert [token.as_dict() for token in document.tokens] == before
	formula_runs = [run for run in output._element.findall('.//w:r', NAMESPACES)
		if ''.join(node.text or '' for node in run.findall('w:t', NAMESPACES)) == formula]
	assert len(formula_runs) == 8
	for run in formula_runs:
		assert run.find('w:rPr/w:color', NAMESPACES).get(docx.oxml.ns.qn('w:val')) == '007A00'
		assert run.find('w:rPr/w:rFonts', NAMESPACES).get(
			docx.oxml.ns.qn('w:ascii')) == 'Roboto Mono'
	linked = output._element.find('.//w:hyperlink/w:r', NAMESPACES)
	assert linked.find('w:rPr/w:u', NAMESPACES) is not None
	compact = [paragraph for paragraph in output.paragraphs
		if paragraph.text.startswith(('Enter ', 'Continuation', 'Explain:'))
		or paragraph.text == formula]
	assert len(compact) == 5
	for paragraph in compact:
		assert paragraph.paragraph_format.space_before.pt == 0
		assert paragraph.paragraph_format.space_after.pt == 2
		assert paragraph.paragraph_format.line_spacing == 1
	body = next(paragraph for paragraph in output.paragraphs if paragraph.text.startswith('Body '))
	assert body.paragraph_format.space_after is None
	assert body.style.paragraph_format.space_after.pt == 6
	assert round(body.style.paragraph_format.line_spacing, 2) == 1.08


#============================================
def test_justification_keeps_deliberate_definition_lines_left_aligned() -> None:
	"""Hard-break definition lines must not acquire stretched justification spacing."""
	source = '# Definitions\n\n'
	source += 'This naturally wrapping paragraph explains a calculation in enough detail that it '
	source += 'will clearly occupy more than two ordinary body lines in the rendered document. '
	source += 'It should remain comfortable to read as a justified narrative paragraph.\n\n'
	source += "**Here's the breakdown:**\\\n"
	source += 'y-hat = predicted exam score\\\n'
	source += 'b0 = intercept\\\n'
	source += 'b1 x = linear term\\\n'
	source += 'x = hours of sleep\n'
	document = tutorial_lib.parser.parse_source(source, 'definitions.md')
	document.content_token_indices = list(range(len(document.tokens)))
	output = docx.Document()
	tutorial_lib.docx_styles.configure(output)
	tutorial_lib.docx_renderer.render_tokens(output, document)
	narrative = next(paragraph for paragraph in output.paragraphs
		if paragraph.text.startswith('This naturally wrapping'))
	definitions = next(paragraph for paragraph in output.paragraphs
		if paragraph.text.startswith("Here's the breakdown:"))
	assert narrative.alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
	assert definitions.alignment is None
	assert definitions.text.count('\n') == 4


#============================================
def test_formula_continuation_becomes_a_copyable_display_without_rewriting_prose() -> None:
	"""An intentional formula line survives Markdown soft-break normalization."""
	source = '# Formula display\n\n## Workbook\n\n### Sheet: Data\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n\n'
	source += '## Tutorial\n\n### Part: Analysis\n\n#### Step: Calculate\n\n'
	source += '* In **E11**, calculate the result using this formula:\n  [`=E6^2 / E5^2`](enter:E11)\n'
	document = tutorial_lib.compiler.compile_source(source)
	assert not document.has_errors
	output = docx.Document()
	tutorial_lib.docx_styles.configure(output)
	tutorial_lib.docx_renderer.render_tokens(output, document)
	intro = next(paragraph for paragraph in output.paragraphs
		if paragraph.text.startswith('In E11,'))
	display = next(paragraph for paragraph in output.paragraphs if paragraph.text == '=E6^2 / E5^2')
	assert intro.text == 'In E11, calculate the result using this formula:'
	assert display._p.find('w:pPr/w:shd', NAMESPACES) is not None
	assert display._p.find('w:pPr/w:pBdr', NAMESPACES) is not None
	assert display.runs[0].font.color.rgb == docx.shared.RGBColor.from_string('007A00')


#============================================
def test_formula_display_keeps_its_justified_introduction_without_trapping_following_prose() -> None:
	"""A displayed formula stays with its long explanation, not the next paragraph."""
	source = '# Formula display\n\n## Workbook\n\n### Sheet: Data\n\n'
	source += '| A | B |\n| --- | --- |\n| 1 | 2 |\n\n'
	source += '## Tutorial\n\n### Part: Analysis\n\n#### Step: Calculate\n\n'
	source += '* In **E11**, carefully calculate the result using this formula after reviewing '
	source += 'the evidence from every measurement and comparing the numerator with the denominator. '
	source += 'This explanation is deliberately long enough to wrap naturally across several lines '\
		'in the rendered document:\n  [`=E6^2 / E5^2`](enter:E11)\n\n'
	source += 'Then explain what the result means in the context of the experiment.\n'
	document = tutorial_lib.compiler.compile_source(source)
	assert not document.has_errors
	output = docx.Document()
	tutorial_lib.docx_styles.configure(output)
	tutorial_lib.docx_renderer.render_tokens(output, document)
	intro = next(paragraph for paragraph in output.paragraphs
		if paragraph.text.startswith('In E11, carefully calculate'))
	display = next(paragraph for paragraph in output.paragraphs if paragraph.text == '=E6^2 / E5^2')
	following = next(paragraph for paragraph in output.paragraphs
		if paragraph.text.startswith('Then explain'))
	assert intro.alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
	assert intro.paragraph_format.keep_with_next is True
	assert not display.paragraph_format.keep_with_next
	assert not following.paragraph_format.keep_with_next


#============================================
def test_literal_action_payloads_keep_copy_text_without_formula_style() -> None:
	"""JSON and Label scalars keep operators, quotes and explicit code formatting."""
	formula = '=A2*B2*3'
	value = '*literal* _text_ <br> a\\b'
	cases = [
		('Enter', json.dumps(formula), formula, False),
		('Enter', '`' + json.dumps(formula) + '`', formula, True),
		('Enter', json.dumps(value), value, False),
		('Replace', json.dumps(formula), formula, False),
		('Label', formula, formula, False),
		('Label', '`' + formula + '`', formula, True),
	]
	for kind, payload, expected_value, code in cases:
		source = '# Literal\n\n## Workbook\n\n### Sheet: Data\n\n'
		source += '| A | B |\n| --- | --- |\n| 1 | 2 |\n\n'
		source += '## Tutorial\n\n### Part: Analysis\n\n#### Step: Calculate\n\n'
		if kind == 'Replace':
			source += '- **Enter C2:** before\n'
		source += f'- **{kind} C2:** ' + payload
		document = tutorial_lib.compiler.compile_source(source)
		assert not document.has_errors
		assert document.actions[-1].scalar.kind == 'text'
		assert document.actions[-1].scalar.value == expected_value
		assert not document.formula_checks
		output = docx.Document()
		tutorial_lib.docx_styles.configure(output)
		tutorial_lib.docx_renderer.render_tokens(output, document)
		paragraph = output.paragraphs[-1]
		expected_text = payload[1:-1] if code else payload
		assert paragraph.text == f'{kind} C2: ' + expected_text
		for run in paragraph.runs:
			assert str(run.font.color.rgb) != '007A00'
		payload_run = paragraph.runs[-1]
		assert payload_run.font.name == ('Roboto Mono' if code else None)


#============================================
def test_standard_styles_and_direct_fonts() -> None:
	"""Paragraph fonts stay editable through styles while role fonts remain explicit."""
	source = "# Title\n\n### Overview\n\n#### Detail\n\n##### Evidence\n\n###### Deep\n\n"
	source += "Body [link](https://example.org) and `=SUM(A1:B1)`\n\n"
	source += "> Quoted evidence\n\n```sheets\n=SUM(A1:B1)\n```\n\n"
	source += "| Note | Formula |\n| --- | --- |\n| [Compact](https://example.org) | `=SUM(A1:B1)` |\n"
	document = tutorial_lib.model.Document(pathlib.Path("styles.md"), source)
	document.tokens = markdown_it.MarkdownIt("commonmark").enable("table").parse(source)
	document.content_token_indices = list(range(len(document.tokens)))
	output = docx.Document()
	tutorial_lib.docx_styles.configure(output)
	tutorial_lib.docx_renderer.render_tokens(output, document)
	standard = {"Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"}
	for name in standard:
		style = output.styles[name]
		assert style.font.name == "PT Sans"
		fonts = style._element.rPr.rFonts
		assert not any("Theme" in key or "theme" in key for key in fonts.attrib)
	assert [p.style.name for p in output.paragraphs[:5]] == [
		"Title", "Heading 1", "Heading 2", "Heading 3", "Heading 3",
	]
	assert all(p.style.name in standard for p in output.paragraphs)
	assert all(run.font.name is None for p in output.paragraphs[:5] for run in p.runs)
	step_heading = output.paragraphs[2]
	body_paragraph = output.paragraphs[5]
	assert step_heading.style.name == "Heading 2"
	assert step_heading.style.paragraph_format.left_indent < body_paragraph.style.paragraph_format.left_indent
	assert output.paragraphs[5].runs[0].font.name is None
	assert output.sections[0].footer.paragraphs[0].style.name == "Normal"
	table_link = output.tables[0].cell(1, 0)._tc.find(".//w:hyperlink/w:r", NAMESPACES)
	assert table_link.find("w:rPr/w:rFonts", NAMESPACES).get(
		docx.oxml.ns.qn("w:ascii")) == "PT Sans Narrow"
	assert output.tables[0].cell(1, 1).paragraphs[0].runs[0].font.name == "Roboto Mono"
	formula = output.paragraphs[-1]
	assert formula.runs[0].font.name == "Roboto Mono"
	assert formula.paragraph_format.keep_together is True
	assert formula._p.find("w:pPr/w:shd", NAMESPACES) is not None
	assert formula._p.find("w:pPr/w:pBdr", NAMESPACES) is not None
	assert output.paragraphs[-2].paragraph_format.left_indent is not None
	runs = output._element.findall(".//w:r", NAMESPACES)
	assert not any(run.find("w:rPr/w:rStyle", NAMESPACES) is not None for run in runs)
	link = output._element.find(".//w:hyperlink/w:r", NAMESPACES)
	assert link.find("w:rPr/w:u", NAMESPACES) is not None
	assert link.find("w:rPr/w:rFonts", NAMESPACES) is None


#============================================
def test_editable_formula_links_and_independent_lists(tmp_path: pathlib.Path) -> None:
	"""Formula copy text and links must survive without rasterizing tutorial content."""
	formula = '=IF(B2>0,"A B",SUM($C$2:C8))'
	source = "# Native tutorial\n\n### Part: First\n\n#### Step: Calculate\n\n"
	source += "3. **Enter D2:** `" + formula + "`\n"
	source += "   - **See:** [calculate](#step-calculate), `B2`, H<sub>2</sub>O<br>next\n\n"
	source += "   - Another child can move independently.\n\n"
	source += "```sheets\n" + formula + "\n```\n\n"
	source += "7. **Choose:** Format > Number\n"
	document = tutorial_lib.model.Document(tmp_path / "input.md", source, title="Native tutorial")
	document.tokens = markdown_it.MarkdownIt("commonmark", {"html": True}).parse(source)
	document.content_token_indices = list(range(len(document.tokens)))
	span = tutorial_lib.model.SourceSpan(str(document.path), 1, 1)
	reference = tutorial_lib.model.CellReference("Data", 2, 2, 2, 2, "B2")
	worksheet = tutorial_lib.model.Worksheet("Data", span)
	worksheet.google = tutorial_lib.model.GoogleMapping("source", "workbook123", "99")
	document.worksheets.append(worksheet)
	for index, token in enumerate(document.tokens):
		if token.type == "heading_open" and token.tag == "h4":
			heading = tutorial_lib.model.Heading(4, "Calculate", "Step 1: Calculate",
				"step-calculate", span, index, kind="step", number="1", prefix_length=6)
			token.meta["heading"] = heading
		if token.children:
			for child in token.children:
				if child.type == "code_inline" and child.content == "B2":
					child.meta["reference"] = reference
	output = tmp_path / "native.docx"
	tutorial_lib.docx_renderer.write_document(document, output)
	with zipfile.ZipFile(output) as archive:
		document_xml = defusedxml.ElementTree.fromstring(archive.read("word/document.xml"))
		numbering = defusedxml.ElementTree.fromstring(archive.read("word/numbering.xml"))
		footer_xml = defusedxml.ElementTree.fromstring(archive.read("word/footer1.xml"))
		relations = archive.read("word/_rels/document.xml.rels").decode("utf-8")
	texts = [node.text for node in document_xml.findall(".//w:t", NAMESPACES)]
	assert texts.count(formula) == 2
	assert "Step 1: Calculate" in "".join(texts)
	assert "workbook123/edit#gid=99&amp;range=B2" in relations
	assert document_xml.find(".//w:hyperlink[@w:anchor]", NAMESPACES) is not None
	assert document_xml.find(".//w:bookmarkStart", NAMESPACES) is not None
	page_field = footer_xml.find(".//w:fldSimple[@w:instr='PAGE']", NAMESPACES)
	assert page_field is not None
	assert document_xml.find(".//w:vertAlign[@w:val='subscript']", NAMESPACES) is not None
	starts = numbering.findall(".//w:start", NAMESPACES)
	assert {"3", "7"}.issubset({node.get("{" + NAMESPACES["w"] + "}val") for node in starts})
	assert document_xml.find(".//w:numPr", NAMESPACES) is not None
	cell_run = next(run for run in document_xml.findall(".//w:r", NAMESPACES)
		if run.find("w:t", NAMESPACES) is not None
		and run.find("w:t", NAMESPACES).text == "B2")
	assert cell_run.find("w:rPr/w:b", NAMESPACES) is not None
	assert cell_run.find("w:rPr/w:rFonts", NAMESPACES).get(
		docx.oxml.ns.qn("w:ascii")) == "Roboto Mono"
	paragraphs = {
		"".join(node.text or "" for node in paragraph.findall(".//w:t", NAMESPACES)): paragraph
		for paragraph in document_xml.findall(".//w:body/w:p", NAMESPACES)
	}
	parent = paragraphs["Enter D2: " + formula]
	child = paragraphs["See: calculate, B2, H2Onext"]
	heading = paragraphs["Step 1: Calculate"]
	assert heading.find(".//w:rFonts", NAMESPACES) is None
	assert parent.find("w:pPr/w:keepNext", NAMESPACES) is not None
	assert child.find("w:pPr/w:keepNext", NAMESPACES) is None
	assert paragraphs["Choose: Format > Number"].find("w:pPr/w:keepNext", NAMESPACES) is None


#============================================
def test_semantic_bindings_preserve_numbered_caffeine_instruction_text() -> None:
	"""Bindings keep their authored paragraph text, heading emphasis and GFM numbering."""
	source = "# Caffeine model\n\n## Workbook\n\n### Sheet: Data\n\n"
	source += "| Sleep | Score |\n| --- | --- |\n| 6 | 91 |\n\n## Tutorial\n\n"
	source += "### **pArT:** **Build** the predictor\n\n"
	source += "#### Step: **Square** the predictor\n\n"
	source += "1. Add a squared predictor column.\n"
	source += "2. Set its heading to [Sleep hours squared](label:C1).\n"
	source += "3. Enter [=A2^2](enter:C2) for the first student.\n"
	source += "4. Copy [C2:C26](fill:C2:C26), then select [A1:B26](select:A1:B26).\n\n"
	source += "See the [original class template](https://example.org/template).\n"
	document = tutorial_lib.compiler.compile_source(source)
	assert not document.has_errors
	assert [action.kind for action in document.actions] == ["label", "enter", "fill", "select"]
	output = docx.Document()
	tutorial_lib.docx_styles.configure(output)
	tutorial_lib.docx_renderer.render_tokens(output, document)
	paragraph_text = [paragraph.text for paragraph in output.paragraphs]
	assert paragraph_text == [
		"Caffeine model", "Part 1: Build the predictor", "Step 1: Square the predictor",
		"Add a squared predictor column.", "Set its heading to Sleep hours squared.",
		"Enter =A2^2 for the first student.",
		"Copy C2:C26, then select A1:B26.",
		"See the original class template.",
	]
	assert any(run.text == "Build" and run.bold for run in output.paragraphs[1].runs)
	assert any(run.text == "Square" and run.bold for run in output.paragraphs[2].runs)
	assert output.paragraphs[4].runs[1].font.color.rgb is None
	formula_run = next(run for run in output.paragraphs[5].runs if run.text == "=A2^2")
	assert str(formula_run.font.color.rgb) == "007A00"
	assert [paragraph.style.name for paragraph in output.paragraphs[3:7]] == [
		"Normal", "Normal", "Normal", "Normal",
	]
	assert [paragraph._p.find("w:pPr/w:numPr", NAMESPACES) is not None
		for paragraph in output.paragraphs[3:7]] == [True, True, True, True]
	assert output.paragraphs[-1]._p.find(".//w:hyperlink", NAMESPACES) is not None
	relationships = output.part.rels
	assert not any("label:" in relation.target_ref or "enter:" in relation.target_ref
		or "fill:" in relation.target_ref or "select:" in relation.target_ref
		for relation in relationships.values())
	assert all(paragraph.paragraph_format.space_after.pt == 2
		for paragraph in output.paragraphs[3:7])
	assert output.styles["Heading 1"].font.size.pt == 18
	assert output.styles["Heading 2"].font.size.pt == 15
	assert output.styles["Heading 3"].font.size.pt == 12


#============================================
def test_result_labels_stay_with_adjacent_entries(tmp_path: pathlib.Path) -> None:
	"""Keep adjacent result pairs without joining chains or crossing block boundaries."""
	source = "# Results\n\n## Workbook\n\n### Sheet: Data\n\n"
	source += "| x | y |\n| --- | --- |\n| 1 | 2 |\n\n"
	source += "## Tutorial\n\n### Part: Analysis\n\n#### Step: Calculate\n\n"
	source += "- **Label C1:** Summary\n"
	source += "- **Label C2:** Slope\n"
	source += "- **Enter D2:** `=SUM(A2:B2)`\n"
	source += "- **Label C3:** Intercept\n\n"
	source += "Explain what this result means.\n\n"
	source += "- **Enter D3:** `=SUM(A2:B2)`\n"
	source += "- **Label C4:** Earlier label\n"
	source += "- **Label C5:** Interpretation\n"
	source += "- **Instruction:** In cell `D5`, write what this result means.\n\n"
	source += "- **Label C6:** Before heading\n\n"
	source += "#### Step: Interpret\n\n"
	source += "- **Instruction:** In cell `D6`, explain the result.\n\n"
	source += "- **Label C7:** Before prose\n\n"
	source += "Consider the evidence.\n\n"
	source += "- **Instruction:** In cell `D7`, explain the evidence.\n\n"
	source += "- Parent\n  - **Label C8:** Nested label\n"
	source += "- **Instruction:** In cell `D8`, explain the nested result.\n\n"
	source += "- **Label C9:** Before table\n\n"
	source += "| Note |\n| --- |\n| Evidence |\n\n"
	source += "- **Instruction:** In cell `D9`, explain the table.\n\n"
	source += "- **Label C10:** Before plain instruction\n"
	source += "- Instruction: In cell `D10`, explain the result.\n"
	source += "- **Label C11:** Before later bold label\n"
	source += "- Read this **Instruction:** carefully.\n"
	source += "\nClosing context.\n\nFinal line.\n"
	document = tutorial_lib.compiler.compile_source(source, tmp_path / "results.md")
	assert not document.has_errors
	output = tmp_path / "results.docx"
	tutorial_lib.docx_renderer.write_document(document, output)
	with zipfile.ZipFile(output) as archive:
		document_xml = defusedxml.ElementTree.fromstring(archive.read("word/document.xml"))
	kept = [
		"".join(node.text or "" for node in paragraph.findall(".//w:t", NAMESPACES))
		for paragraph in document_xml.findall(".//w:body/w:p", NAMESPACES)
		if paragraph.find("w:pPr/w:keepNext", NAMESPACES) is not None
	]
	assert "Label C2: Slope" in kept
	assert "Label C1: Summary" not in kept
	assert "Label C3: Intercept" not in kept
	assert "Enter D2: =SUM(A2:B2)" not in kept
	assert "Label C5: Interpretation" in kept
	assert "Instruction: In cell D5, write what this result means." not in kept
	assert not set(kept).intersection({
		"Label C4: Earlier label", "Label C6: Before heading", "Label C7: Before prose",
		"Label C8: Nested label", "Label C9: Before table",
		"Label C10: Before plain instruction", "Label C11: Before later bold label",
	})


#============================================
def test_semantic_adjacent_label_stays_with_direct_or_nested_entry(tmp_path: pathlib.Path) -> None:
	"""Only adjacent sibling labels pair with same-row entries, including first nested formulas."""
	base = "# Pagination\n\n## Workbook\n\n### Sheet: Data\n\n| A | B | C | D | E |\n"
	base += "| --- | --- | --- | --- | --- |\n| 1 | 2 | 3 | 4 | 5 |\n\n## Tutorial\n\n"
	base += "### Part: Analysis\n\n#### Step: Prepare\n\n"
	cases = {
		"direct": "1. Label [D13](label:D13)\n2. Enter [=1](enter:E13)\n",
		"nested": "1. Label [D13](label:D13)\n2. In cell E13, enter:\n"
			"   - [=1](enter:E13)\n   - Explanation text.\n",
		"different_row": "1. Label [D13](label:D13)\n2. Enter [=1](enter:E14)\n",
		"heading_barrier": "1. Label [D13](label:D13)\n\n#### Step: Next\n\n"
			"1. Enter [=1](enter:E13)\n",
	}
	for name, body in cases.items():
		source = base + body + "\nClosing context.\n\nFinal line.\n"
		document = tutorial_lib.compiler.compile_source(source, tmp_path / (name + ".md"))
		assert not document.has_errors
		path = tmp_path / (name + ".docx")
		tutorial_lib.docx_renderer.write_document(document, path)
		with zipfile.ZipFile(path) as archive:
			xml = defusedxml.ElementTree.fromstring(archive.read("word/document.xml"))
		paragraphs = {
			"".join(node.text or "" for node in item.findall(".//w:t", NAMESPACES)): item
			for item in xml.findall(".//w:body/w:p", NAMESPACES)
		}
		is_kept = paragraphs["Label D13"].find("w:pPr/w:keepNext", NAMESPACES) is not None
		assert is_kept is (name in ("direct", "nested"))
		if name == "nested":
			parent_keep = paragraphs["In cell E13, enter:"].find("w:pPr/w:keepNext", NAMESPACES)
			assert parent_keep is not None
			assert paragraphs["=1"].find("w:pPr/w:keepNext", NAMESPACES) is None
			assert paragraphs["Explanation text."].find("w:pPr/w:keepNext", NAMESPACES) is None


#============================================
def test_final_body_paragraph_stays_with_its_predecessor(tmp_path: pathlib.Path) -> None:
	"""A short closing paragraph cannot be stranded by itself at a page break."""
	source = "# Closing\n\n## Workbook\n\n### Sheet: Data\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n\n"
	source += "## Tutorial\n\n### Part: Analysis\n\n#### Step: Finish\n\n"
	source += "Review the result before moving on.\n\n"
	source += "**_Ask the instructor for the next tutorial link._**\n"
	document = tutorial_lib.compiler.compile_source(source, tmp_path / "closing.md")
	assert not document.has_errors
	path = tmp_path / "closing.docx"
	tutorial_lib.docx_renderer.write_document(document, path)
	with zipfile.ZipFile(path) as archive:
		xml = defusedxml.ElementTree.fromstring(archive.read("word/document.xml"))
	paragraphs = {
		"".join(node.text or "" for node in item.findall(".//w:t", NAMESPACES)): item
		for item in xml.findall(".//w:body/w:p", NAMESPACES)
	}
	assert paragraphs["Review the result before moving on."].find(
		"w:pPr/w:keepNext", NAMESPACES) is not None
	assert paragraphs["Ask the instructor for the next tutorial link."].find(
		"w:pPr/w:keepNext", NAMESPACES) is None


#============================================
@pytest.mark.skipif(not PIL.features.check("webp"), reason="Pillow requires WebP support")
def test_local_webp_embeds_lossless_png_without_changing_source(tmp_path: pathlib.Path) -> None:
	"""WebP screenshots retain pixels, alpha, proportions, alt text and original asset identity."""
	asset = tmp_path / "screenshot.webp"
	image = PIL.Image.new("RGBA", (4, 2), (12, 34, 56, 255))
	image.putpixel((1, 0), (200, 100, 50, 128))
	image.save(asset, format="WEBP", lossless=True, exact=True)
	original = asset.read_bytes()
	source = "# Image\n\n## Workbook\n\n### Sheet: Data\n\n| A |\n| --- |\n| 1 |\n\n"
	source += "## Tutorial\n\n### Part: Review\n\n#### Step: Observe\n\n"
	source += "![Screenshot evidence](screenshot.webp)\n"
	document = tutorial_lib.compiler.compile_source(source, tmp_path / "images.md")
	assert not document.has_errors
	assert document.assets == [asset.resolve()]
	output = tmp_path / "images.docx"
	tutorial_lib.docx_renderer.write_document(document, output)
	reopened = docx.Document(output)
	shape = reopened.inline_shapes[0]
	assert shape.width == 2 * shape.height
	assert shape._inline.docPr.get("descr") == "Screenshot evidence"
	assert shape._inline.docPr.get("title") == "Screenshot evidence"
	with zipfile.ZipFile(output) as archive:
		media = [name for name in archive.namelist() if name.startswith("word/media/")]
		assert len(media) == 1
		with PIL.Image.open(io.BytesIO(archive.read(media[0]))) as embedded:
			assert embedded.format == "PNG"
			assert embedded.size == image.size
			assert embedded.convert("RGBA").tobytes() == image.tobytes()
	assert asset.read_bytes() == original
	assert document.source == source
	assert list(tmp_path.glob("*.png")) == []
	for invalid in (b"not an image", b"RIFF\x00\x00\x00\x00WEBPbroken"):
		asset.write_bytes(invalid)
		rejected = tutorial_lib.compiler.compile_source(source, tmp_path / "images.md")
		assert any(item.category == "image-invalid" for item in rejected.diagnostics)


#============================================
def test_tall_images_fit_native_printable_page_without_cropping(tmp_path: pathlib.Path) -> None:
	"""Tall screenshots scale into the body while keeping source pixels and alt text."""
	for extension in ("png", "webp"):
		asset = tmp_path / f"tall.{extension}"
		PIL.Image.new("RGB", (300, 1800), (12, 34, 56)).save(asset)
		original = asset.read_bytes()
		source = f"![Tall evidence]({asset.name})\n"
		document = tutorial_lib.parser.parse_source(source, tmp_path / "images.md")
		document.content_token_indices = list(range(len(document.tokens)))
		output = docx.Document()
		tutorial_lib.docx_styles.configure(output)
		section = output.sections[0]
		# Exercise native dimensions rather than assuming the default paper size.
		section.page_height = docx.shared.Inches(8)
		section.top_margin = docx.shared.Inches(1)
		section.bottom_margin = docx.shared.Inches(1)
		tutorial_lib.docx_renderer.render_tokens(output, document)
		path = tmp_path / f"{extension}.docx"
		output.save(path)
		reopened = docx.Document(path)
		shape = reopened.inline_shapes[0]
		assert 0 < shape.height < section.page_height - section.top_margin - section.bottom_margin
		assert 0 < shape.width <= section.page_width - section.left_margin - section.right_margin
		assert shape.height / shape.width == pytest.approx(6)
		assert shape._inline.docPr.get("descr") == "Tall evidence"
		with zipfile.ZipFile(path) as archive:
			media = next(name for name in archive.namelist() if name.startswith("word/media/"))
			with PIL.Image.open(io.BytesIO(archive.read(media))) as embedded:
				assert embedded.size == (300, 1800)
		assert asset.read_bytes() == original
		assert document.source == source


#============================================
def test_native_tables_images_and_menu_segments(tmp_path: pathlib.Path) -> None:
	"""Tables and images retain editable content, embedded data and authored order."""
	png = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aS1sAAAAASUVORK5CYII="
	(tmp_path / "pixel.png").write_bytes(base64.b64decode(png))
	source = "Before [linked before ![Local alt](pixel.png) linked after](https://example.org) after\n\n"
	source += "| Group | Value |\n| --- | --- |\n| native | 12 |\n\n"
	source += "![Data alt](data:image/png;base64," + png + ")\n\n"
	source += "- **Choose:** Format > Number > Percent\n"
	document = tutorial_lib.model.Document(tmp_path / "images.md", source)
	parser = markdown_it.MarkdownIt("commonmark").enable("table")
	document.tokens = parser.parse(source)
	document.content_token_indices = list(range(len(document.tokens)))
	span = tutorial_lib.model.SourceSpan(str(document.path), 1, 1)
	for index, token in enumerate(document.tokens):
		if token.type == "inline" and token.content.startswith("**Choose:"):
			token.meta["action"] = tutorial_lib.model.Action(
				"choose", span, index, "Format > Number > Percent",
				menu=["Format", "Number", "Percent"],
			)
	output = tmp_path / "images.docx"
	tutorial_lib.docx_renderer.write_document(document, output)
	with zipfile.ZipFile(output) as archive:
		document_xml = defusedxml.ElementTree.fromstring(archive.read("word/document.xml"))
		relations = archive.read("word/_rels/document.xml.rels").decode("utf-8")
		assert any(name.startswith("word/media/") for name in archive.namelist())
	texts = "".join(node.text or "" for node in document_xml.findall(".//w:t", NAMESPACES))
	assert "Before linked before  linked after after" in texts
	link = document_xml.find(".//w:hyperlink", NAMESPACES)
	linked_content = []
	for run in link.findall("w:r", NAMESPACES):
		if run.find("w:drawing", NAMESPACES) is not None:
			linked_content.append("image")
		else:
			linked_content.append("".join(node.text or "" for node in run.findall("w:t", NAMESPACES)))
	assert linked_content == ["linked before ", "image", " linked after"]
	assert 'Target="https://example.org"' in relations
	assert "native12" in texts
	assert "Format > Number > Percent" in texts
	menu_run = next(run for run in document_xml.findall(".//w:r", NAMESPACES)
		if "Format >" in "".join(run.itertext()))
	assert menu_run.find("w:rPr/w:b", NAMESPACES) is not None
	assert menu_run.find("w:rPr/w:rFonts", NAMESPACES).get(
		docx.oxml.ns.qn("w:ascii")) == "Roboto Mono"
	assert len(document_xml.findall(".//w:tbl", NAMESPACES)) == 1
	rows = document_xml.findall(".//w:tbl/w:tr", NAMESPACES)
	assert all(row.find("w:trPr/w:cantSplit", NAMESPACES) is not None for row in rows)
	assert rows[0].find("w:trPr/w:tblHeader", NAMESPACES) is not None
	assert rows[1].find("w:trPr/w:tblHeader", NAMESPACES) is None
	assert document_xml.find(".//w:tbl/w:tr/w:tc/w:p/w:pPr/w:keepNext", NAMESPACES) is None
	assert len(document_xml.findall(".//w:drawing", NAMESPACES)) == 2
