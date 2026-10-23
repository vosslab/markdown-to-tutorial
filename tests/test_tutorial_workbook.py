"""Protect workbook seed integrity and fail-closed one-time import behavior."""

import pathlib
import zipfile

import pytest
import openpyxl

import tutorial_lib.model
import tutorial_lib.compiler
import tutorial_lib.workbook
import tutorial_lib.formula_io


#============================================
def test_seed_types_and_student_targets_roundtrip(tmp_path: pathlib.Path) -> None:
	"""Student actions cannot leak answers or coerce literal and empty text seeds."""
	model = tutorial_lib.model
	span = model.SourceSpan("source.md", 1, 1)
	sheet = model.Worksheet("Data", span)
	values = [model.Scalar("text", "=literal"), model.Scalar("text", ""),
		model.Scalar("blank", None), model.Scalar("number", 1.2345678901234567),
		model.Scalar("boolean", False), model.Scalar("formula", '=IF(A1="x",2,3)')]
	for column, scalar in enumerate(values, 1):
		sheet.cells[(4, column)] = model.CellSeed(4, column, scalar, span)
	sheet.entry_cells = {(4, 1), (4, 3), (100, 8)}
	document = model.Document(pathlib.Path("source.md"), "", worksheets=[sheet])
	document.actions = [model.Action("Enter", span, 0, "secret answer")]
	path = tmp_path / "student.xlsx"
	tutorial_lib.workbook.write_workbook(document, path)
	tutorial_lib.workbook.verify_workbook(document, path)
	book = openpyxl.load_workbook(path, data_only=False)
	assert book["Data"]["A4"].data_type == "s"
	assert book["Data"]["H100"].value is None
	assert book["Data"]["C4"].fill.fill_type == "solid"
	assert book["Data"]["A4"].fill.fill_type == "solid"
	assert book["Data"].freeze_panes is None
	assert book["Data"].column_dimensions["H"].width == 12
	with zipfile.ZipFile(path) as archive:
		assert b"secret answer" not in b"".join(archive.read(n) for n in archive.namelist())
	book.close()


#============================================
def test_semantic_workbook_roles_style_regions_without_changing_content(
	tmp_path: pathlib.Path,
) -> None:
	"""Role styling distinguishes authored regions and preserves every workbook value."""
	model = tutorial_lib.model
	span = model.SourceSpan("source.md", 1, 1)
	sheet = model.Worksheet("Data", span)
	sheet.regions.append(model.TableRegion(1, 1, 3, 2, span))
	for address, row, column, value, kind in (
		("A1", 1, 1, "Sample", "text"), ("B1", 1, 2, "Value", "text"),
		("A2", 2, 1, 1.25, "number"), ("B2", 2, 2, 2.5, "number"),
		("D1", 1, 4, "Mean", "text"), ("E1", 1, 5, "Calculated mean", "text"),
		("D2", 2, 4, "", "text"), ("E2", 2, 5, "", "text"),
		("D3", 3, 4, "Conclusion", "text"),
	):
		sheet.cells[(row, column)] = model.CellSeed(
			row, column, model.Scalar(kind, value), span)
	sheet.entry_cells.update({(2, 5), (3, 4)})
	sheet.role_regions.extend([
		model.RoleRegion("results", model.CellReference("Data", 1, 4, 2, 5, "D1:E2"), span),
		model.RoleRegion("interpretation",
			model.CellReference("Data", 3, 4, 3, 5, "D3:E3"), span),
	])
	document = model.Document(pathlib.Path("source.md"), "", worksheets=[sheet])
	path = tmp_path / "roles.xlsx"
	tutorial_lib.workbook.write_workbook(document, path)
	tutorial_lib.workbook.verify_workbook(document, path)
	book = openpyxl.load_workbook(path, data_only=False)
	output = book["Data"]
	assert output["A1"].fill.fgColor.rgb.endswith("24527A")
	assert output["A2"].fill.fgColor.rgb.endswith("DCEAF7")
	assert output["D1"].fill.fgColor.rgb.endswith("A64B00")
	assert output["E2"].fill.fgColor.rgb.endswith("FCE6CF")
	assert output["D3"].fill.fgColor.rgb.endswith("EAE2F5")
	assert output["E3"].fill.fgColor.rgb.endswith("EAE2F5")
	assert output["E2"].value is None
	assert output["E2"].border.bottom.style == "medium"
	assert output["C1"].fill.fill_type is None
	assert output["A1"].font.name == "PT Sans"
	assert output["A1"].font.color.rgb.endswith("FFFFFF")
	book.close()


#============================================
def test_modern_formula_storage_preserves_logical_formulas(tmp_path: pathlib.Path) -> None:
	"""XLSX prefixes enable engine recognition without rewriting text or unknown calls."""
	model = tutorial_lib.model
	span = model.SourceSpan("source.md", 1, 1)
	sheet = model.Worksheet("Data", span)
	formulas = [
		'=  IF(TRUE, STDEV.S(A1:A2), "NORM.S.DIST( and _xlfn.STDEV.S(")',
		'=_xlfn.NORM.DIST(1,0,1,TRUE)',
		'=SUM(1,  2)+MY.UNKNOWN(3)+_xlfn.UNKNOWN(4)',
	]
	for column, formula in enumerate(formulas, 1):
		sheet.cells[(3, column)] = model.CellSeed(
			3, column, model.Scalar("formula", formula), span)
	document = model.Document(pathlib.Path("source.md"), "", worksheets=[sheet])
	path = tmp_path / "modern.xlsx"
	tutorial_lib.workbook.write_workbook(document, path)
	tutorial_lib.workbook.verify_workbook(document, path)
	book = openpyxl.load_workbook(path, data_only=False)
	assert book["Data"]["A3"].value == formulas[0].replace(
		" STDEV.S(A1:A2)", " _xlfn.STDEV.S(A1:A2)")
	assert book["Data"]["B3"].value == formulas[1]
	assert book["Data"]["C3"].value == formulas[2]
	book.close()
	draft = tutorial_lib.workbook.import_workbook(path)
	assert formulas[0] in draft
	assert '=NORM.DIST(1,0,1,TRUE)' in draft
	assert '_xlfn.UNKNOWN(4)' in draft
	for formula in formulas:
		stored = tutorial_lib.formula_io.serialize_formula(formula)
		assert tutorial_lib.formula_io.serialize_formula(stored) == stored
		assert tutorial_lib.formula_io.logical_formula(stored) == (
			tutorial_lib.formula_io.logical_formula(formula))


#============================================
def test_row_one_table_freezes_header(tmp_path: pathlib.Path) -> None:
	"""Only a table whose header starts in row one gets a frozen top row."""
	model = tutorial_lib.model
	span = model.SourceSpan("source.md", 1, 1)
	sheet = model.Worksheet("Data", span)
	sheet.regions.append(model.TableRegion(1, 1, 2, 3, span))
	sheet.cells[(1, 1)] = model.CellSeed(1, 1, model.Scalar("text", "Name"), span)
	sheet.cells[(1, 2)] = model.CellSeed(1, 2, model.Scalar("text", "Result"), span)
	sheet.entry_cells.add((2, 3))
	document = model.Document(pathlib.Path("source.md"), "", worksheets=[sheet])
	path = tmp_path / "header.xlsx"
	tutorial_lib.workbook.write_workbook(document, path)
	book = openpyxl.load_workbook(path)
	assert book["Data"].freeze_panes == "A2"
	assert book["Data"].column_dimensions["C"].width == 12
	book.close()


#============================================
def test_blank_label_targets_reserve_space_for_all_alternatives(tmp_path: pathlib.Path) -> None:
	"""Labels remain readable beside entered results without revealing student answers."""
	source = "# Analysis\n\n## Workbook\n\n### Sheet: Data\n\n"
	source += "| Observation |\n| --- |\n| 1 |\n\n## Tutorial\n\n"
	source += "### Part: Summarize\n\n- **Worksheet:** Data\n\n#### Step: Calculate\n\n"
	source += "- **Label D4:** Sample standard deviation\n- **Enter E4:** `=1+1`\n"
	source += "- **Alternative: Short**\n  - **Label H11:** p-value\n"
	source += "- **Alternative: Full**\n"
	source += "  - **Label H11:** Two-tailed p-value (quadratic term differs from zero)\n"
	source += "- **Enter I11:** `=1/2`\n"
	source += '- **Enter E5:** `=IF(E4<0,"Reject the Null Hypothesis",'
	source += '"Fail to Reject the Null Hypothesis")`\n'
	source += '- **Enter E6:** Record your interpretation\n'
	document = tutorial_lib.compiler.compile_source(source)
	assert not document.has_errors
	path = tmp_path / "labels.xlsx"
	tutorial_lib.workbook.write_workbook(document, path)
	tutorial_lib.workbook.verify_workbook(document, path)
	book = openpyxl.load_workbook(path)
	sheet = book["Data"]
	assert all(sheet[address].value is None
		for address in ("D4", "E4", "E5", "E6", "H11", "I11"))
	assert sheet.column_dimensions["D"].width >= len("Sample standard deviation") + 2
	assert sheet.column_dimensions["E"].width >= len("Fail to Reject the Null Hypothesis") + 2
	assert 12 < sheet.column_dimensions["H"].width <= 48
	assert sheet.row_dimensions[11].height >= 30
	assert all(sheet[address].alignment.wrap_text for address in ("D4", "E4", "H11", "I11"))
	# Entered values retain the generated layout instead of relying on text overflow.
	sheet["H11"] = "Two-tailed p-value (quadratic term differs from zero)"
	sheet["I11"] = 0.125
	book.save(path)
	book.close()
	book = openpyxl.load_workbook(path)
	assert book["Data"]["H11"].alignment.wrap_text
	assert book["Data"].row_dimensions[11].height >= 30
	book.close()


#============================================
def test_import_preserves_sparse_addresses_and_literal_types(tmp_path: pathlib.Path) -> None:
	"""Sparse sheet imports remain bounded and quote ambiguous text and formulas faithfully."""
	book = openpyxl.Workbook()
	sheet = book.active
	sheet.title = "Data"
	sheet["C8"] = "001"
	sheet["D8"] = ""
	sheet["D8"].data_type = "s"
	sheet["E8"] = "literal &amp; <br> `ticks` \\| and\nnewline"
	sheet["F1000000"] = '=IF(C8="a|b",1,0)'
	book.create_sheet("Empty")
	path = tmp_path / "sparse.xlsx"
	book.save(path)
	book.close()
	draft = tutorial_lib.workbook.import_workbook(path)
	assert "#### At: C8" in draft
	assert "#### At: F1000000" in draft
	assert '` "001" `' in draft and '` "" `' in draft
	assert r'=IF(C8="a\|b",1,0)' in draft
	assert "### Sheet: Empty\n\n|  |\n| --- |" in draft
	assert len(draft) < 1500
	# Finish the intentionally incomplete scaffold to exercise the shared scalar parser.
	source = draft.replace("Write the tutorial instructions here.",
		"### Part: Review\n\n- **Worksheet:** Data\n\n"
		"#### Step: Inspect\n\n- **Compare:** Inspect the imported cells.")
	document = tutorial_lib.compiler.compile_source(source)
	assert not document.has_errors
	rebuilt = tmp_path / "rebuilt.xlsx"
	tutorial_lib.workbook.write_workbook(document, rebuilt)
	tutorial_lib.workbook.verify_workbook(document, rebuilt)
	book = openpyxl.load_workbook(rebuilt, data_only=False)
	assert book["Data"]["E8"].value == "literal &amp; <br> `ticks` \\| and\nnewline"
	assert book["Data"]["F1000000"].value == '=IF(C8="a|b",1,0)'
	book.close()


#============================================
@pytest.mark.parametrize("feature", ["merge", "date", "error", "epoch", "iteration"])
def test_import_blocks_semantic_loss(tmp_path: pathlib.Path, feature: str) -> None:
	"""A draft must not silently discard workbook and cell semantics."""
	book = openpyxl.Workbook()
	sheet = book.active
	if feature == "merge":
		sheet.merge_cells("A1:B1")
	elif feature == "date":
		sheet["A1"] = 45000
		sheet["A1"].number_format = "yyyy-mm-dd"
	elif feature == "epoch":
		book.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
		sheet["A1"] = "=DATE(2026,1,1)"
		sheet["B1"] = 44561
	elif feature == "iteration":
		book.calculation = openpyxl.workbook.properties.CalcProperties(
			iterate=True, iterateCount=1000, iterateDelta=0.0001)
		sheet["A1"] = "=(A1+1)/2"
	else:
		sheet["A1"] = "#DIV/0!"
	path = tmp_path / "unsupported.xlsx"
	book.save(path)
	book.close()
	message = "import blocked"
	if feature == "epoch":
		message = "1904 date epoch"
	elif feature == "iteration":
		message = "iterative calculation"
	with pytest.raises(ValueError, match=message):
		tutorial_lib.workbook.import_workbook(path)
