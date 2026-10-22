#!/usr/bin/env python3
"""Execute final tutorial states in LibreOffice Calc and compare independent expectations."""

import json
import shutil
import pathlib
import hashlib
import argparse
import datetime
import tempfile
import subprocess

import openpyxl

import tutorial_lib.compiler
import tutorial_lib.formula_io
import sheets_validation_lib


#============================================
def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(description=__doc__)
	parser.add_argument("tutorials", nargs="+", type=pathlib.Path)
	parser.add_argument("-e", "--expectations", dest="expectations", type=pathlib.Path,
		required=True)
	parser.add_argument("-o", "--output", dest="output", type=pathlib.Path,
		default=pathlib.Path("tutorial_output/validation/libreoffice"))
	args = parser.parse_args()
	return args


#============================================
def save_json(path: pathlib.Path, payload: object) -> None:
	path.write_text(json.dumps(payload, indent=2, allow_nan=False) + "\n", encoding="utf-8")


#============================================
def find_engine() -> str:
	executable = shutil.which("libreoffice") or shutil.which("soffice")
	macos_binary = pathlib.Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
	if executable is None and macos_binary.is_file():
		executable = str(macos_binary)
	if executable is None:
		raise RuntimeError("Local calculation validation requires LibreOffice Calc.")
	return executable


#============================================
def write_case(case: dict, path: pathlib.Path) -> dict:
	"""Apply typed seeds and authored replacements in order without evaluating formulas."""
	workbook = openpyxl.Workbook()
	workbook.remove(workbook.active)
	for name in case["sheets"]:
		workbook.create_sheet(name)
	final_writes = {}
	for write in case["seeds"] + case["writes"]:
		cell = workbook[write["sheet"]][write["address"]]
		value = write["value"]
		cell.value = next(iter(value.values())) if value else None
		if "formulaValue" in value:
			cell.value = tutorial_lib.formula_io.serialize_formula(value["formulaValue"])
		# Explicit text prevents a literal equals prefix from becoming an XLSX formula.
		if "stringValue" in value:
			cell.data_type = "s"
		final_writes[write["sheet"] + "!" + write["address"]] = write
	workbook.calculation.fullCalcOnLoad = True
	workbook.calculation.forceFullCalc = True
	workbook.save(path)
	workbook.close()
	return final_writes


#============================================
def inspect_results(path: pathlib.Path, final_writes: dict, expected: dict) -> dict:
	"""Read actual engine caches, rejecting errors and absent final formula results."""
	workbook = openpyxl.load_workbook(path, data_only=True)
	cells, errors = {}, []
	for worksheet in workbook:
		for row in worksheet:
			for cell in row:
				if cell.value is None:
					continue
				address = worksheet.title + "!" + cell.coordinate
				value = cell.value
				if cell.data_type == "e":
					effective = {"errorValue": {"type": "CALC_ERROR", "message": value}}
					errors.append({"address": address, "error": value})
				elif type(value) is bool:
					effective = {"boolValue": value}
				elif type(value) in (int, float):
					effective = {"numberValue": value}
				else:
					effective = {"stringValue": value}
				cells[address] = {"effectiveValue": effective}
	workbook.close()
	for address, write in final_writes.items():
		if "formulaValue" in write["value"] and address not in cells:
			errors.append({"address": address, "error": "MISSING_CALCULATED_CACHE"})
	for error in errors:
		write = final_writes.get(error["address"])
		if write is not None:
			error.update(source_line=write["line"], authored_value=write["value"])
	comparisons = sheets_validation_lib.compare_expectations(cells, expected)
	for comparison in comparisons:
		write = final_writes[comparison["address"]]
		comparison.update(source_line=write["line"], authored_value=write["value"])
	result = {"status": "failed" if errors or any(item["status"] == "failed"
		for item in comparisons) else "passed", "errors": errors, "comparisons": comparisons,
		"cells": cells}
	return result


#============================================
def run_case(executable: str, case: dict, expected: dict, destination: pathlib.Path) -> dict:
	"""Use a private profile and distinct input/output directories for Calc conversion."""
	destination.mkdir(parents=True, exist_ok=False)
	with tempfile.TemporaryDirectory(prefix="tutorial-calc-") as temporary:
		staging = pathlib.Path(temporary)
		source = staging / "scenario.xlsx"
		final_writes = write_case(case, source)
		formulas = {address for address, write in final_writes.items()
			if "formulaValue" in write["value"]}
		if formulas != set(expected):
			raise ValueError(f"{case['tutorial']} {case['scenario']}: expectation mapping "
				f"missing={sorted(formulas - set(expected))}, extra={sorted(set(expected) - formulas)}")
		command = [executable, f"-env:UserInstallation={(staging / 'profile').as_uri()}",
			"--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir",
			str(destination.resolve()), str(source)]
		process = subprocess.run(command, capture_output=True, text=True, timeout=120)
	converted = destination / source.name
	result = {"tutorial": case["tutorial"], "scenario": case["scenario"],
		"branches": case["branches"], "process_returncode": process.returncode,
		"stdout": process.stdout, "stderr": process.stderr, "workbook": str(converted)}
	if process.returncode != 0 or not converted.is_file():
		result.update(status="failed", errors=[{"error": "ENGINE_CONVERSION_FAILED"}],
			comparisons=[])
	else:
		result.update(inspect_results(converted, final_writes, expected))
	return result


#============================================
def main() -> int:
	args = parse_args()
	executable = find_engine()
	version = subprocess.run([executable, "--version"], capture_output=True, text=True,
		check=True, timeout=30).stdout.strip()
	payload = json.loads(args.expectations.read_text(encoding="utf-8"))["tutorials"]
	cases, hashes = [], {}
	for path in args.tutorials:
		cases.extend(sheets_validation_lib.prepare_cases(tutorial_lib.compiler.compile_document(path)))
		hashes[str(path)] = hashlib.sha256(path.read_bytes()).hexdigest()
	if {pathlib.Path(case["tutorial"]).name for case in cases} != set(payload):
		raise ValueError("Tutorial selection must match the expectation file's tutorial names.")
	args.output.mkdir(parents=True, exist_ok=True)
	# Each invocation keeps separate artifacts; old workbooks can never satisfy a new run.
	run_directory = pathlib.Path(tempfile.mkdtemp(prefix="run-", dir=args.output))
	report = {"status": "running", "engine": version, "engine_executable": executable,
		"created_at": datetime.datetime.now(datetime.timezone.utc).isoformat(),
		"scope": "final-state calculation only; intermediate errors and Google Sheets unverified",
		"source_sha256": hashes, "expectations_sha256": hashlib.sha256(
			args.expectations.read_bytes()).hexdigest(), "results": []}
	report["validator_sha256"] = hashlib.sha256(pathlib.Path(__file__).read_bytes()).hexdigest()
	report["formula_io_sha256"] = hashlib.sha256(
		pathlib.Path(tutorial_lib.formula_io.__file__).read_bytes()).hexdigest()
	report_path = run_directory / "report.json"
	save_json(report_path, report)
	for index, case in enumerate(cases):
		item = payload[pathlib.Path(case["tutorial"]).name]
		scenario_values = item["scenarios"]
		if set(scenario_values) - {entry["scenario"] for entry in cases
			if entry["tutorial"] == case["tutorial"]}:
			raise ValueError(f"Unknown expectation scenario: {case['tutorial']}")
		expected = item["common"] | scenario_values.get(case["scenario"], {})
		result = run_case(executable, case, expected, run_directory / f"case-{index + 1:02d}")
		report["results"].append(result)
		save_json(report_path, report)
		print(f"{case['tutorial']} {case['scenario']}: {result['status']}")
	report["status"] = "passed" if all(item["status"] == "passed"
		for item in report["results"]) else "failed"
	report["scenario_count"] = len(cases)
	report["comparison_count"] = sum(len(item["comparisons"]) for item in report["results"])
	save_json(report_path, report)
	print(f"LibreOffice Calc {report['status']}. Report: {report_path}")
	return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
	raise SystemExit(main())
