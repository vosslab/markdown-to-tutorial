"""Render Markdown inline content into editable Word runs and links."""

import io
import re
import base64
import hashlib
import pathlib
import urllib.parse

import docx.oxml
import docx.oxml.ns
import docx.shared
import docx.opc.constants

import tutorial_lib.docx_styles


#============================================
def bookmark_name(slug: str) -> str:
	"""Use a stable Word-safe bookmark name for a compiled heading slug."""
	# A digest avoids truncation collisions and Word's identifier length limit.
	name = "heading_" + hashlib.sha256(slug.encode("utf-8")).hexdigest()[:24]
	return name


#============================================
def reference_url(document: object, reference: object) -> str | None:
	worksheet = next(sheet for sheet in document.worksheets if sheet.name == reference.sheet)
	if worksheet.google is None or worksheet.google.gid is None:
		return None
	mapping = worksheet.google
	fragment = urllib.parse.urlencode({"gid": mapping.gid, "range": reference.address})
	url = f"https://docs.google.com/spreadsheets/d/{mapping.spreadsheet_id}/edit#{fragment}"
	return url


#============================================
def link_element(paragraph: object, href: str) -> object:
	link = docx.oxml.OxmlElement("w:hyperlink")
	if href.startswith("#"):
		link.set(docx.oxml.ns.qn("w:anchor"), bookmark_name(urllib.parse.unquote(href[1:])))
	else:
		relation = paragraph.part.relate_to(
			href, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True,
		)
		link.set(docx.oxml.ns.qn("r:id"), relation)
	paragraph._p.append(link)
	return link


#============================================
def add_text(paragraph: object, text: str, state: dict, role: str | None = None,
	parent: object | None = None) -> object:
	"""Write a native run and optionally place it inside a hyperlink."""
	run = paragraph.add_run(text)
	run.bold = state["bold"]
	run.italic = state["italic"]
	run.font.subscript = state["sub"]
	run.font.superscript = state["sup"]
	tutorial_lib.docx_styles.format_run(
		run, role, linked=parent is not None, font_name=state["font"],
	)
	if parent is not None:
		parent.append(run._r)
	return run


#============================================
def add_image(paragraph: object, token: object, document: object,
	parent: object | None = None) -> None:
	"""Embed local or data-URI PNG/JPEG content at its original aspect ratio."""
	source = token.attrGet("src")
	if source.startswith("data:"):
		header, payload = source.split(",", 1)
		if header not in ("data:image/png;base64", "data:image/jpeg;base64"):
			raise ValueError("DOCX images require PNG/JPEG base64 data URIs")
		stream = io.BytesIO(base64.b64decode(payload, validate=True))
	else:
		path = pathlib.Path(urllib.parse.unquote(source))
		if not path.is_absolute():
			path = document.path.parent / path
		stream = str(path)
	run = paragraph.add_run()
	shape = run.add_picture(stream)
	# Keep the image at the current position within an authored link.
	if parent is not None:
		parent.append(run._r)
	maximum = docx.shared.Inches(tutorial_lib.docx_styles.IMAGE_WIDTH)
	if shape.width > maximum:
		height = round(shape.height * maximum / shape.width)
		shape.width = maximum
		shape.height = height
	shape._inline.docPr.set("descr", token.content)
	shape._inline.docPr.set("title", token.content)


#============================================
def has_binding(tokens: list) -> bool:
	"""Return whether an inline token sequence carries semantic link bindings."""
	return any("binding_action" in token.meta for token in tokens)


#============================================
def render(paragraph: object, tokens: list, document: object, action: object = None,
	font_name: str | None = None) -> None:
	"""Preserve inline emphasis, HTML scripts, formula bytes, images and links."""
	state = {"bold": False, "italic": False, "sub": False, "sup": False, "font": font_name}
	links = []
	menu_rendered = False
	for token in tokens:
		kind = token.type
		if kind in ("strong_open", "strong_close"):
			state["bold"] = kind.endswith("open")
		elif kind in ("em_open", "em_close"):
			state["italic"] = kind.endswith("open")
		elif kind == "link_open":
			if "binding_action" in token.meta:
				# Semantic links carry spreadsheet meaning, not a navigable URI.
				links.append(None)
			else:
				links.append(link_element(paragraph, token.attrGet("href")))
		elif kind == "link_close":
			links.pop()
		elif kind == "image":
			add_image(paragraph, token, document, parent=links[-1] if links else None)
		elif kind in ("softbreak", "hardbreak"):
			add_text(paragraph, " " if kind == "softbreak" else "\n", state,
				parent=links[-1] if links else None)
		elif kind == "html_inline":
			tag = token.content.lower().strip()
			if re.fullmatch(r"<br\s*/?>", tag):
				add_text(paragraph, "\n", state, parent=links[-1] if links else None)
			elif tag in ("<sub>", "</sub>", "<sup>", "</sup>"):
				state["sub" if "sub" in tag else "sup"] = not tag.startswith("</")
			else:
				raise ValueError(f"Unsupported DOCX inline HTML: {token.content}")
		elif kind in ("text", "code_inline"):
			role = token.meta.get("role", "code" if kind == "code_inline" else None)
			text = token.content
			parent = links[-1] if links else None
			if (action is not None and action.kind.lower() == "choose"
				and not has_binding(tokens) and not state["bold"]):
				if text.strip() and not menu_rendered:
					add_text(paragraph, " " + " > ".join(action.menu), state, "menu")
					menu_rendered = True
				continue
			if (action is not None and action.target is not None and state["bold"]
				and not has_binding(tokens)):
				address = action.target.address
				if address in text and text.startswith(action.kind.capitalize() + " "):
					before, address, after = text.partition(address)
					add_text(paragraph, before, state, parent=parent)
					url = reference_url(document, action.target)
					if url is not None and parent is None:
						parent = link_element(paragraph, url)
					add_text(paragraph, address, state, "cell", parent)
					add_text(paragraph, after, state, parent=links[-1] if links else None)
					continue
			if "reference" in token.meta:
				role = "cell"
				url = reference_url(document, token.meta["reference"])
				if url is not None and parent is None:
					parent = link_element(paragraph, url)
			add_text(paragraph, text, state, role, parent)
		else:
			raise ValueError(f"Unsupported DOCX inline token: {kind}")
