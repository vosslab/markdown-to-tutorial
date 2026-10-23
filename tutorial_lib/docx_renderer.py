"""Native DOCX export of the compiled tutorial, excluding workbook declarations."""

import pathlib
import copy

import docx
import docx.oxml
import docx.oxml.ns
import docx.shared

import tutorial_lib.model
import tutorial_lib.docx_inline
import tutorial_lib.docx_styles
import tutorial_lib.docx_numbering


#============================================
def add_bookmark(paragraph: object, slug: str, identifier: int) -> None:
	start = docx.oxml.OxmlElement("w:bookmarkStart")
	start.set(docx.oxml.ns.qn("w:id"), str(identifier))
	start.set(docx.oxml.ns.qn("w:name"), tutorial_lib.docx_inline.bookmark_name(slug))
	end = docx.oxml.OxmlElement("w:bookmarkEnd")
	end.set(docx.oxml.ns.qn("w:id"), str(identifier))
	paragraph._p.insert(0, start)
	paragraph._p.append(end)


#============================================
def add_heading(output: object, token: object, inline: object, document: object,
	identifier: int) -> None:
	"""Preserve authored heading runs and use native outline styles."""
	heading = token.meta.get("heading")
	if heading is None:
		level = int(token.tag[1:])
		style = "Title" if level == 1 else f"Heading {max(1, min(level - 2, 3))}"
		paragraph = output.add_paragraph(style=style)
		tutorial_lib.docx_inline.render(paragraph, inline.children or [], document)
		return
	styles = {"part": "Heading 1", "step": "Heading 2",
		"substep": "Heading 3", "title": "Title"}
	style = styles.get(heading.kind, f"Heading {max(1, min(heading.level - 2, 3))}")
	if heading.level == 1:
		style = "Title"
	paragraph = output.add_paragraph(style=style)
	children = list(inline.children or [])
	if heading.prefix_length:
		remaining = heading.prefix_length
		for child_index, child in enumerate(children):
			if child.type not in ("text", "code_inline") or not child.content:
				continue
			children[child_index] = copy.copy(child)
			removed = min(remaining, len(child.content))
			children[child_index].content = child.content[removed:]
			remaining -= removed
			if remaining == 0:
				break
		if remaining:
			raise ValueError("Heading prefix length exceeds authored heading text")
		prefix = heading.display[:-len(heading.title)] if heading.title else heading.display
		paragraph.add_run(prefix)
	tutorial_lib.docx_inline.render(paragraph, children, document)
	add_bookmark(paragraph, heading.slug, identifier)


#============================================
def keep_label_with_entry(paragraph: object, document: tutorial_lib.model.Document,
	position: int, label: tutorial_lib.model.Action | None,
	actions: dict[int, tutorial_lib.model.Action]) -> None:
	"""Keep a label with a direct sibling entry or its first nested formula."""
	if label is None or label.kind != "label" or label.target is None:
		return
	indices = document.content_token_indices
	position += 1
	if position + 4 >= len(indices):
		return
	shape = [document.tokens[indices[index]].type for index in range(position, position + 5)]
	if shape != ["paragraph_close", "list_item_close", "list_item_open",
		"paragraph_open", "inline"]:
		return
	inline_index = indices[position + 4]
	if document.tokens[indices[position - 1]].level != document.tokens[inline_index].level:
		return

	def action_at(index: int) -> tutorial_lib.model.Action | None:
		inline = document.tokens[index]
		action = inline.meta.get("action", actions.get(index))
		if action is None:
			action = next((child.meta.get("binding_action") for child in inline.children or []
				if child.meta.get("binding_action") is not None), None)
		return action

	entry = action_at(inline_index)
	if entry is None:
		# Preserve the authored Label -> bold Instruction pairing.
		children = document.tokens[inline_index].children or []
		start = 0
		while start < len(children) and children[start].type == "text" and not children[start].content:
			start += 1
		opening = children[start:start + 3]
		if ([child.type for child in opening] == ["strong_open", "text", "strong_close"]
			and opening[1].content.strip() == "Instruction:"):
			paragraph.paragraph_format.keep_with_next = True
			return
		# A numbered instruction can place its formula in its first nested item.
		nested_position = position + 5
		if nested_position + 4 >= len(indices):
			return
		nested_shape = [document.tokens[indices[index]].type
			for index in range(nested_position, nested_position + 5)]
		if nested_shape[0] != "paragraph_close" or nested_shape[1] not in (
			"bullet_list_open", "ordered_list_open") or nested_shape[2:] != [
				"list_item_open", "paragraph_open", "inline"]:
			return
		nested_inline = indices[nested_position + 4]
		if document.tokens[nested_inline].level != document.tokens[inline_index].level + 2:
			return
		entry = action_at(nested_inline)
	if entry is None or entry.kind not in ("enter", "replace") or entry.target is None:
		return

	left, right = label.target, entry.target
	if (left.start_row == left.end_row == right.start_row == right.end_row
		and left.start_column == left.end_column
		and right.start_column == right.end_column == left.start_column + 1
		and left.sheet == right.sheet):
		paragraph.paragraph_format.keep_with_next = True


#============================================
def trim_terminal_hardbreak(children: list) -> list:
	"""Drop a terminal list hardbreak that would add an empty line to the next item."""
	result = list(children)
	index = len(result) - 1
	closing = {"strong_close", "em_close", "s_close", "link_close"}
	while index >= 0 and (result[index].type in closing
		or result[index].type == "text" and not result[index].content):
		index -= 1
	if index >= 0 and result[index].type == "hardbreak":
		result.pop(index)
	return result


#============================================
def display_formula_break(children: list) -> tuple[list, list, tutorial_lib.model.Action | None]:
	"""Split an authored formula continuation into a compact, editable formula display.

	A soft break normally remains ordinary prose.  This narrow pattern preserves the
	author's intentional separate formula line when its final content is one semantic
	formula binding, as in a list item ending with ``[=A2^2](enter:C2)``.
	"""
	for position, child in enumerate(children):
		if child.type != "softbreak":
			continue
		suffix = children[position + 1:]
		action = next((item.meta.get("binding_action") for item in suffix
			if item.meta.get("binding_action") is not None), None)
		if action is None or action.scalar is None or action.scalar.kind != "formula":
			continue
		meaningful = [item for item in suffix if item.type not in (
			"link_open", "link_close", "strong_open", "strong_close", "em_open", "em_close")
			and not (item.type == "text" and not item.content.strip())]
		if len(meaningful) == 1 and meaningful[0].content == action.scalar.value:
			return children[:position], suffix, action
	return children, [], None


#============================================
def add_formula_display(output: object, children: list, document: tutorial_lib.model.Document,
	lists: list, items: list) -> None:
	"""Render a standalone instruction formula without turning ordinary prose into cards."""
	paragraph = output.add_paragraph(style="Normal")
	tutorial_lib.docx_styles.format_formula(paragraph)
	if items:
		paragraph.paragraph_format.left_indent = docx.shared.Inches(
			tutorial_lib.docx_styles.LIST_INDENT_INCHES * len(lists))
	tutorial_lib.docx_inline.render(paragraph, children, document)


#============================================
def formula_only(children: list) -> bool:
	"""Keep a dedicated one-line formula left aligned, even outside a fence."""
	meaningful = [child for child in children if child.type not in (
		"link_open", "link_close", "strong_open", "strong_close", "em_open", "em_close")
		and not (child.type == "text" and not child.content.strip())]
	return len(meaningful) == 1 and meaningful[0].meta.get("role") == "formula"


#============================================
def render_tokens(output: object, document: tutorial_lib.model.Document) -> None:
	"""Walk selected block tokens while retaining list and table structure."""
	lists = []
	items = []
	paragraph = None
	table = None
	row = None
	cell = None
	column = 0
	quote_depth = 0
	action = None
	headings = {item.token_index: item for item in document.headings}
	actions = {item.token_index: item for item in document.actions}
	indices = document.content_token_indices
	for position, index in enumerate(indices):
		token = document.tokens[index]
		kind = token.type
		if kind == "heading_open":
			# The compiler also exposes the heading map for exporter consumers.
			if "heading" not in token.meta and index in headings:
				# Never mutate the compiled token model.
				token = copy.copy(token)
				token.meta = dict(token.meta, heading=headings[index])
			add_heading(output, token, document.tokens[indices[position + 1]], document, index)
		elif kind == "inline":
			if position and document.tokens[indices[position - 1]].type == "heading_open":
				continue
			if paragraph is None:
				raise ValueError("Inline tutorial content has no containing paragraph")
			bound = tutorial_lib.docx_inline.has_binding(token.children or [])
			action = None if bound else token.meta.get("action", actions.get(index))
			pagination_action = action
			if bound:
				pagination_action = next((child.meta.get("binding_action")
					for child in token.children or [] if child.meta.get("binding_action") is not None), None)
			font_name = None
			if cell is not None:
				font_name = tutorial_lib.docx_styles.NARROW_FONT
			children = token.children or []
			if items:
				children = trim_terminal_hardbreak(children)
			first_children, formula_children, formula_action = display_formula_break(children)
			tutorial_lib.docx_inline.render(
				paragraph, first_children, document, action, font_name=font_name,
			)
			if cell is None and not formula_only(first_children):
				indent = tutorial_lib.docx_styles.BODY_INDENT_INCHES
				if quote_depth:
					indent += 0.25
				if items:
					indent += tutorial_lib.docx_styles.LIST_INDENT_INCHES * len(lists)
				width = tutorial_lib.docx_styles.USABLE_PAGE_WIDTH_INCHES - indent
				tutorial_lib.docx_styles.justify_narrative(paragraph, width)
			if formula_action is not None:
				# The narrative introduces this display; keep that relationship on the
				# introductory paragraph so the formula can still flow with its own
				# following prose.
				paragraph.paragraph_format.keep_with_next = True
				add_formula_display(output, formula_children, document, lists, items)
			if cell is None:
				keep_label_with_entry(paragraph, document, position, pagination_action, actions)
		elif kind in ("bullet_list_open", "ordered_list_open"):
			# Keep a parent label with its first child, without locking the whole list.
			if lists and items and position:
				previous = document.tokens[indices[position - 1]]
				if previous.type == "paragraph_close":
					output.paragraphs[-1].paragraph_format.keep_with_next = True
			start = int(token.attrGet("start") or 1)
			number = tutorial_lib.docx_numbering.create_list(
				output, kind == "ordered_list_open", start, len(lists),
			)
			lists.append(number)
		elif kind in ("bullet_list_close", "ordered_list_close"):
			lists.pop()
		elif kind == "list_item_open":
			items.append(False)
			action = token.meta.get("action", actions.get(index))
		elif kind == "list_item_close":
			items.pop()
			action = None
		elif kind == "paragraph_open":
			action = token.meta.get("action", actions.get(index, action))
			paragraph = cell.paragraphs[0] if cell is not None else output.add_paragraph(style="Normal")
			if quote_depth:
				paragraph.paragraph_format.left_indent = docx.shared.Inches(0.25)
			if items and cell is None:
				tutorial_lib.docx_styles.format_list(paragraph)
				if not items[-1]:
					tutorial_lib.docx_numbering.apply_list(paragraph, lists[-1])
					items[-1] = True
				else:
					paragraph.paragraph_format.left_indent = docx.shared.Inches(
						tutorial_lib.docx_styles.LIST_INDENT_INCHES * len(lists),
					)
		elif kind == "paragraph_close":
			paragraph = None
		elif kind in ("fence", "code_block"):
			if output.paragraphs:
				output.paragraphs[-1].paragraph_format.keep_with_next = True
			paragraph = output.add_paragraph(style="Normal")
			tutorial_lib.docx_styles.format_formula(paragraph)
			if items:
				tutorial_lib.docx_styles.format_list(paragraph)
			# Markdown fences carry a parser-added final newline, outside the formula.
			run = paragraph.add_run(token.content.removesuffix("\n"))
			tutorial_lib.docx_styles.format_run(run, token.meta.get("role", "code"))
			paragraph = None
		elif kind == "blockquote_open":
			quote_depth += 1
		elif kind == "blockquote_close":
			quote_depth -= 1
		elif kind == "table_open":
			# Count the header's cells, without inferring spreadsheet semantics.
			columns = 0
			for following in indices[position + 1:]:
				if document.tokens[following].type == "tr_close":
					break
				if document.tokens[following].type in ("th_open", "td_open"):
					columns += 1
			table = output.add_table(rows=0, cols=columns)
			table.style = "Table Grid"
		elif kind == "tr_open":
			row = table.add_row()
			properties = row._tr.get_or_add_trPr()
			# A row that exceeds a page can still flow; ordinary rows stay intact.
			properties.append(docx.oxml.OxmlElement("w:cantSplit"))
			if len(table.rows) == 1:
				properties.append(docx.oxml.OxmlElement("w:tblHeader"))
			elif len(table.rows) % 2 == 1:
				tutorial_lib.docx_styles.format_table_stripe(row)
			column = 0
		elif kind in ("th_open", "td_open"):
			cell = row.cells[column]
			paragraph = cell.paragraphs[0]
			paragraph.style = "Normal"
			paragraph.paragraph_format.left_indent = docx.shared.Inches(0)
			if kind == "th_open":
				tutorial_lib.docx_styles.shade(
					cell._tc.get_or_add_tcPr(), tutorial_lib.docx_styles.TABLE_HEADER_FILL,
				)
		elif kind in ("th_close", "td_close"):
			if kind == "th_close":
				tutorial_lib.docx_styles.format_table_header(cell)
			cell = None
			paragraph = None
			column += 1
		elif kind == "table_close":
			table = None
		elif kind not in ("heading_close", "tr_close", "thead_open", "thead_close",
			"tbody_open", "tbody_close"):
			raise ValueError(f"Unsupported DOCX block token: {kind}")
	body_blocks = [element for element in output._element.body
		if element.tag in (docx.oxml.ns.qn("w:p"), docx.oxml.ns.qn("w:tbl"))]
	if (len(body_blocks) > 1 and body_blocks[-1].tag == docx.oxml.ns.qn("w:p")
		and body_blocks[-2].tag == docx.oxml.ns.qn("w:p")):
		# Keep a short closing paragraph from becoming a page by itself.
		penultimate = output.paragraphs[-2]
		properties = penultimate._p.find(docx.oxml.ns.qn("w:pPr"))
		is_formula_display = (properties is not None
			and properties.find(docx.oxml.ns.qn("w:shd")) is not None)
		if not is_formula_display:
			penultimate.paragraph_format.keep_with_next = True


#============================================
def write_document(document: tutorial_lib.model.Document, output_path: str | pathlib.Path) -> None:
	"""Write a compiled tutorial as editable DOCX without altering the model."""
	if document.has_errors:
		raise ValueError("Cannot render a tutorial with semantic errors")
	output = docx.Document()
	tutorial_lib.docx_styles.configure(output)
	output.core_properties.title = document.title
	output.core_properties.subject = "Spreadsheet tutorial"
	render_tokens(output, document)
	output.save(output_path)
