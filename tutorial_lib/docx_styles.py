"""Central style definitions for native, editable tutorial documents."""

import docx.enum.text
import docx.oxml
import docx.oxml.ns
import docx.shared
import pathlib

from PIL import ImageFont


BODY_FONT = "PT Sans"
NARROW_FONT = "PT Sans Narrow"
CODE_FONT = "Roboto Mono"
BODY_SIZE = 11
CODE_SIZE = 9
PAGE_MARGIN = 0.8
IMAGE_WIDTH = 6.8
LIST_INDENT_TWIPS = 360
LIST_HANGING_TWIPS = 240
LIST_INDENT_INCHES = 0.25
TABLE_HEADER_FILL = "0067CC"
TABLE_STRIPE_FILL = "EEF5FC"
FORMULA_COLOR = "007A00"
FORMULA_FILL = "E7F4EA"
FORMULA_BORDER = "007A00"
BODY_INDENT_INCHES = 0.2
USABLE_PAGE_WIDTH_INCHES = 6.9
FONT_PATHS = (
	pathlib.Path("/System/Library/Fonts/Supplemental/PTSans.ttc"),
	pathlib.Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
)
FONT_CACHE = {}


#============================================
def _set_explicit_font(style: object, font_name: str) -> None:
	"""Set a font family and remove theme selectors that override it in exports."""
	style.font.name = font_name
	r_fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
	for theme_attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
		r_fonts.attrib.pop(docx.oxml.ns.qn(f"w:{theme_attribute}"), None)
	for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
		r_fonts.set(docx.oxml.ns.qn(f"w:{attribute}"), font_name)


#============================================
def format_run(run: object, role: str | None = None, linked: bool = False,
	font_name: str | None = None) -> None:
	"""Apply portable role formatting directly, without custom character styles."""
	if role in ("code", "cell", "menu", "formula"):
		font_name = CODE_FONT
		run.font.size = docx.shared.Pt(CODE_SIZE)
	# Ordinary text inherits its editable paragraph style; only roles override fonts.
	if font_name is not None:
		_set_explicit_font(run, font_name)
	if role in ("cell", "menu"):
		run.bold = True
	if linked:
		run.font.underline = True
		run.font.color.rgb = docx.shared.RGBColor.from_string("174F91")
	if role == "formula":
		run.font.color.rgb = docx.shared.RGBColor.from_string(FORMULA_COLOR)


#============================================
def format_list(paragraph: object) -> None:
	"""Compact every paragraph within a list, including continuation paragraphs."""
	paragraph.paragraph_format.space_before = docx.shared.Pt(0)
	paragraph.paragraph_format.space_after = docx.shared.Pt(2)
	paragraph.paragraph_format.line_spacing = 1.0


#============================================
def format_table_header(cell: object) -> None:
	"""Give authored tables a compact, readable teaching-material header."""
	shade(cell._tc.get_or_add_tcPr(), TABLE_HEADER_FILL)
	for paragraph in cell.paragraphs:
		paragraph.paragraph_format.space_after = docx.shared.Pt(2)
		for run in paragraph.runs:
			run.bold = True
			run.font.color.rgb = docx.shared.RGBColor.from_string("FFFFFF")


#============================================
def format_table_stripe(row: object) -> None:
	"""Use a quiet alternating tint so dense scientific tables remain scannable."""
	for cell in row.cells:
		shade(cell._tc.get_or_add_tcPr(), TABLE_STRIPE_FILL)


#============================================
def estimated_line_count(text: str, width_inches: float, size: int = BODY_SIZE) -> int:
	"""Estimate editable Word line wrapping with installed body-font metrics.

	Word and Google Docs choose their own final line breaks.  This estimate deliberately
	only selects justification for clearly long prose, where a one-line difference does
	not change readability.
	"""
	key = size
	if key not in FONT_CACHE:
		path = next((item for item in FONT_PATHS if item.exists()), None)
		FONT_CACHE[key] = (ImageFont.truetype(str(path), size)
			if path is not None else ImageFont.load_default())
	font = FONT_CACHE[key]
	limit = max(1, width_inches * 72)
	lines = 0
	for source_line in text.splitlines() or [text]:
		words = source_line.split()
		if not words:
			lines += 1
			continue
		current = ""
		for word in words:
			candidate = word if not current else current + " " + word
			if current and font.getlength(candidate) > limit:
				lines += 1
				current = word
			else:
				current = candidate
		lines += 1
	return lines


#============================================
def justify_narrative(paragraph: object, width_inches: float) -> None:
	"""Justify naturally wrapping prose, while preserving deliberate display lines.

	A Markdown hard break becomes a native line break in the editable document.  Those
	breaks are used for equation definitions and compact term-by-term explanations;
	LibreOffice expands the spaces on every such line when the paragraph is justified.
	They are display content rather than prose that merely happened to wrap.
	"""
	if "\n" not in paragraph.text and estimated_line_count(paragraph.text, width_inches) > 2:
		paragraph.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY


#============================================
def shade(element: object, color: str) -> None:
	"""Apply an OOXML background without replacing the native content."""
	fill = docx.oxml.OxmlElement("w:shd")
	fill.set(docx.oxml.ns.qn("w:fill"), color)
	element.append(fill)


#============================================
def add_page_number(paragraph: object) -> None:
	"""Add a native Word PAGE field to a footer paragraph."""
	paragraph.style = "Normal"
	paragraph.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
	run = paragraph.add_run()
	field = docx.oxml.OxmlElement("w:fldSimple")
	field.set(docx.oxml.ns.qn("w:instr"), "PAGE")
	result = docx.oxml.OxmlElement("w:r")
	text = docx.oxml.OxmlElement("w:t")
	text.text = "1"
	result.append(text)
	field.append(result)
	run._r.addnext(field)


#============================================
def configure(document: object) -> None:
	"""Configure the six standard paragraph styles used by tutorial content."""
	for section in document.sections:
		section.top_margin = docx.shared.Inches(PAGE_MARGIN)
		section.bottom_margin = docx.shared.Inches(PAGE_MARGIN)
		section.left_margin = docx.shared.Inches(PAGE_MARGIN)
		section.right_margin = docx.shared.Inches(PAGE_MARGIN)
		add_page_number(section.footer.paragraphs[0])
	styles = document.styles
	normal = styles["Normal"]
	_set_explicit_font(normal, BODY_FONT)
	normal.font.size = docx.shared.Pt(BODY_SIZE)
	normal.paragraph_format.space_after = docx.shared.Pt(6)
	normal.paragraph_format.line_spacing = 1.08
	normal.paragraph_format.left_indent = docx.shared.Inches(BODY_INDENT_INCHES)
	for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
		style = styles[name]
		_set_explicit_font(style, BODY_FONT)
		if name.startswith("Heading"):
			level = int(name[-1])
			size = {1: 18, 2: 15, 3: 12}[level]
			style.font.size = docx.shared.Pt(size)
			style.font.bold = True
			style.paragraph_format.keep_with_next = True
			style.paragraph_format.space_before = docx.shared.Pt(14)
			style.paragraph_format.space_after = docx.shared.Pt(6)
			style.paragraph_format.left_indent = docx.shared.Inches(
				{"Heading 1": 0.1, "Heading 2": 0, "Heading 3": 0.12}[name])
			style.font.color.rgb = docx.shared.RGBColor.from_string(
				{"Heading 1": "003FFF", "Heading 2": "007576", "Heading 3": "B74300"}[name])
	title = styles["Title"]
	title.font.size = docx.shared.Pt(25)
	title.font.bold = True
	title.font.color.rgb = docx.shared.RGBColor.from_string("0067CC")
	title.paragraph_format.space_after = docx.shared.Pt(14)
	title.paragraph_format.left_indent = docx.shared.Inches(0)


#============================================
def format_formula(paragraph: object) -> None:
	"""Keep editable formula blocks visually distinct using direct formatting."""
	paragraph.paragraph_format.left_indent = docx.shared.Inches(BODY_INDENT_INCHES)
	paragraph.paragraph_format.right_indent = docx.shared.Inches(0.12)
	paragraph.paragraph_format.space_before = docx.shared.Pt(5)
	paragraph.paragraph_format.space_after = docx.shared.Pt(8)
	properties = paragraph._p.get_or_add_pPr()
	shade(properties, FORMULA_FILL)
	# Word paragraph borders provide visible padding and a non-color cue.
	borders = docx.oxml.OxmlElement("w:pBdr")
	for edge in ("top", "left", "bottom", "right"):
		border = docx.oxml.OxmlElement(f"w:{edge}")
		for key, value in {"val": "single", "sz": "5", "space": "5", "color": FORMULA_BORDER}.items():
			border.set(docx.oxml.ns.qn(f"w:{key}"), value)
		borders.append(border)
	properties.append(borders)
