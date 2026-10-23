"""Typed student workbook export and deliberately limited XLSX migration."""

# Standard Library
import json
import math
import re
import pathlib
import textwrap
import zipfile

# PIP modules
import defusedxml.ElementTree
import openpyxl
import openpyxl.styles
import openpyxl.formula.tokenizer
import openpyxl.workbook.properties

# Local modules
import tutorial_lib.model
import tutorial_lib.formula_io


MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
MAX_IMPORT_CELLS = 250000

ROLE_STYLES = {
	"source": ("EEF5FC", "0067CC", "FFFFFF"),
	"calculation": ("E6F3F0", "00775F", "FFFFFF"),
	"results": ("F4E9FB", "A719DB", "FFFFFF"),
	"interpretation": ("FDEAF3", "CC0066", "FFFFFF"),
}

FORMULA_ENTRY_STYLE = ("E7F4EA", "007A00")
LABEL_ENTRY_STYLE = ("FFF3E5", "B74300")
SOURCE_STRIPE_FILL = "EEF5FC"
GRIDLINE_COLOR = "C9D8E8"
INTERPRETATION_PANEL_FILL = "FFF7FB"
DEFAULT_NUMBER_FORMAT = "0.000"


#============================================
def _xml_cells(path: pathlib.Path) -> dict[str, dict[str, tuple[str, bool]]]:
	"""Read storage types that openpyxl collapses, notably explicit empty strings."""
	result = {}
	with zipfile.ZipFile(path) as archive:
		for name in archive.namelist():
			if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
				continue
			root = defusedxml.ElementTree.fromstring(archive.read(name))
			cells = {}
			for cell in root.iter(f"{MAIN_NS}c"):
				kind = cell.attrib.get("t", "n")
				cells[cell.attrib["r"]] = (kind, cell.find(f"{MAIN_NS}f") is not None)
			result[name] = cells
	return result


#============================================
def _sheet_parts(path: pathlib.Path) -> list[str]:
	"""Resolve worksheet parts in workbook order rather than guessing sheet IDs."""
	with zipfile.ZipFile(path) as archive:
		root = defusedxml.ElementTree.fromstring(archive.read("xl/workbook.xml"))
		rels = defusedxml.ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
		mapping = {item.attrib["Id"]: item.attrib["Target"] for item in rels}
		parts = []
		for sheet in root.find(f"{MAIN_NS}sheets"):
			rel_id = sheet.attrib[
				"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
			target = mapping[rel_id]
			part = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
			parts.append(part)
	return parts


#============================================
def _logical_cells(path: pathlib.Path) -> tuple[list[str], list[dict]]:
	"""Read logical values without formula caches, retaining text and blank types."""
	storage = _xml_cells(path)
	parts = _sheet_parts(path)
	book = openpyxl.load_workbook(path, data_only=False)
	all_cells = []
	for sheet, part in zip(book.worksheets, parts, strict=True):
		cells = {}
		for address, (kind, formula) in storage[part].items():
			cell = sheet[address]
			value = cell.value
			if formula:
				scalar = tutorial_lib.model.Scalar("formula",
					tutorial_lib.formula_io.logical_formula(value))
			elif kind in ("s", "inlineStr", "str"):
				scalar = tutorial_lib.model.Scalar("text", "" if value is None else value)
			elif value is None:
				scalar = tutorial_lib.model.Scalar("blank", None)
			elif isinstance(value, bool):
				scalar = tutorial_lib.model.Scalar("boolean", value)
			elif isinstance(value, (int, float)):
				if not math.isfinite(value):
					raise ValueError(f"{sheet.title}!{address}: nonfinite number is unsupported")
				scalar = tutorial_lib.model.Scalar("number", value)
			else:
				raise ValueError(f"{sheet.title}!{address}: unsupported cell value {value!r}")
			cells[(cell.row, cell.column)] = scalar
		all_cells.append(cells)
	book.close()
	return book.sheetnames, all_cells


#============================================
def write_workbook(document: tutorial_lib.model.Document, output_path: pathlib.Path | str) -> None:
	"""Write only authored seeds, with distinct styling for empty student targets."""
	if document.has_errors:
		raise ValueError("Cannot export a document with compilation errors")
	if not document.worksheets:
		raise ValueError("A workbook requires at least one worksheet")
	book = openpyxl.Workbook()
	book.remove(book.active)
	book.calculation = openpyxl.workbook.properties.CalcProperties(
		calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode="auto")
	for source in document.worksheets:
		sheet = book.create_sheet(source.name)
		if any(region.start_row == 1 for region in source.regions):
			sheet.freeze_panes = "A2"
		widths = {}
		for (row, column), seed in source.cells.items():
			cell = sheet.cell(row, column)
			scalar = seed.scalar
			if scalar.kind == "number" and not math.isfinite(scalar.value):
				raise ValueError(f"{source.name}!{cell.coordinate}: nonfinite number")
			if scalar.kind == "text" and len(scalar.value) > 32767:
				raise ValueError(f"{source.name}!{cell.coordinate}: text exceeds XLSX 32767 characters")
			cell.value = scalar.value
			if scalar.kind == "formula":
				cell.value = tutorial_lib.formula_io.serialize_formula(scalar.value)
			if scalar.kind == "text":
				# Leading equals signs and explicit empty strings remain literal text.
				cell.data_type = "s"
			if scalar.kind in ("number", "formula", "blank"):
				# Retain the stored value while making ordinary scientific results easy to read.
				cell.number_format = DEFAULT_NUMBER_FORMAT
			cell.font = openpyxl.styles.Font(name="PT Sans", size=10, color="202124")
			cell.alignment = openpyxl.styles.Alignment(vertical="top", wrap_text=True)
			length = max((len(line) for line in str(scalar.value).splitlines()), default=0)
			widths[column] = min(48, max(widths.get(column, 12), length + 2))

		# Style only the usable part of each authored area.  A role range can reserve
		# space for learners, but it must not become a mysterious colored rectangle.
		role_cells = {}
		region_bounds = []
		for region in source.regions:
			bounds = _content_bounds(source, region.start_row, region.start_column,
				region.end_row, region.end_column, include_entries=False)
			if bounds is None:
				continue
			region_bounds.append((bounds, "source", region.start_row))
			for row in range(bounds[0], bounds[2] + 1):
				for column in range(bounds[1], bounds[3] + 1):
					role_cells[(row, column)] = "source"
		for row, column in source.entry_cells:
			role_cells[(row, column)] = "calculation"
		for region in source.role_regions:
			bounds = _content_bounds(source, region.reference.start_row,
				region.reference.start_column, region.reference.end_row,
				region.reference.end_column, include_entries=True)
			# A response panel can intentionally be empty in the starter workbook.
			# It is still a real authored space, unlike unused padding in a data table.
			if region.role == "interpretation":
				bounds = (region.reference.start_row, region.reference.start_column,
					region.reference.end_row, region.reference.end_column)
			if bounds is None:
				continue
			region_bounds.append((bounds, region.role, region.reference.start_row))
			for row in range(bounds[0], bounds[2] + 1):
				for column in range(bounds[1], bounds[3] + 1):
					role_cells[(row, column)] = region.role
		future_headers = _future_headers(document, source, region_bounds)
		for (row, column), role in role_cells.items():
			if role not in ROLE_STYLES:
				raise ValueError(f"Unknown workbook role {role!r} in {source.name}")
			background, header_color, header_text = ROLE_STYLES[role]
			cell = sheet.cell(row, column)
			if role == "source" and _source_stripe(source, row, column):
				cell.fill = openpyxl.styles.PatternFill("solid", fgColor=SOURCE_STRIPE_FILL)
			elif role == "interpretation":
				cell.fill = openpyxl.styles.PatternFill("solid", fgColor=INTERPRETATION_PANEL_FILL)
			elif (row, column) in source.entry_cells:
				cell.fill = openpyxl.styles.PatternFill("solid", fgColor=background)
			if role == "source":
				cell.border = openpyxl.styles.Border(
					bottom=openpyxl.styles.Side(style="thin", color=GRIDLINE_COLOR))
			cell.font = openpyxl.styles.Font(name="PT Sans", size=10, color="202124")
			cell.alignment = openpyxl.styles.Alignment(vertical="top", wrap_text=True)
			seed = source.cells.get((row, column))
			if seed is None or seed.scalar.kind == "blank":
				cell.number_format = DEFAULT_NUMBER_FORMAT
		for bounds, role, _ in region_bounds:
			_apply_region_outline(sheet, bounds[0], bounds[1], bounds[2], bounds[3],
				ROLE_STYLES[role][1])

		# Emphasize supplied headers and planned headers.  A planned header is a
		# genuine learner label task in the first row of an authored area.  Its
		# format prepares the completed table without supplying the label itself.
		header_ranges = [(bounds[0], bounds[1], bounds[3])
			for bounds, role, header_row in region_bounds
			if role == "source" and bounds[0] == header_row]
		for bounds, role, header_row in region_bounds:
			if role == "source" or bounds[0] == bounds[2]:
				continue
			row = bounds[0]
			columns = range(bounds[1], bounds[3] + 1)
			if any(isinstance(sheet.cell(row, column).value, str)
				and sheet.cell(row, column).value.strip() for column in columns):
				header_ranges.append((row, bounds[1], bounds[3]))
		for row, start_column, end_column in future_headers:
			header_ranges.append((row, start_column, end_column))
		future_header_cells = {(row, column) for row, start_column, end_column in future_headers
			for column in range(start_column, end_column + 1)}
		for row, start_column, end_column in _unique_header_ranges(header_ranges):
			for column in range(start_column, end_column + 1):
				cell = sheet.cell(row, column)
				role = role_cells.get((row, column))
				if role is None:
					continue
				_, header_color, header_text = ROLE_STYLES[role]
				cell.fill = openpyxl.styles.PatternFill("solid", fgColor=header_color)
				cell.font = openpyxl.styles.Font(
					name="PT Sans", size=10, bold=True, color=header_text)
				cell.alignment = openpyxl.styles.Alignment(
					vertical="center", horizontal="center", wrap_text=True)
				sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 15, 30)
			if row == 1:
				sheet.freeze_panes = "A2"
		# Authored text estimates space without inserting or evaluating student answers.
		authored_texts = {}
		for action in document.actions:
			if action.kind not in ("label", "enter", "replace"):
				continue
			if action.target.sheet != source.name:
				continue
			row, column = action.target.start_row, action.target.start_column
			for text in _action_texts(action):
				authored_texts.setdefault((row, column), []).append(text)
				length = max((len(line) for line in text.splitlines()), default=0)
				widths[column] = min(48, max(widths.get(column, 12), length + 2))
		for row, column in source.entry_cells:
			# Empty answer columns have no seed from which to infer a useful width.
			widths[column] = max(widths.get(column, 12), 12)
			seed = source.cells.get((row, column))
			if (row, column) in future_header_cells:
				continue
			if (seed is None or seed.scalar.kind == "blank"
				or seed.scalar.kind == "text" and seed.scalar.value == ""):
				cell = sheet.cell(row, column)
				cell.alignment = openpyxl.styles.Alignment(vertical="top", wrap_text=True)
				background, header_color = _entry_style(document, source.name, row, column,
					role_cells.get((row, column), "calculation"))
				cell.fill = openpyxl.styles.PatternFill("solid", fgColor=background)
				cell.border = openpyxl.styles.Border(
					left=cell.border.left, right=cell.border.right, top=cell.border.top,
					bottom=openpyxl.styles.Side(style="thin", color=header_color))
		for column, width in widths.items():
			sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
		# Explicit heights survive typing beside an occupied result cell without auto-fit.
		for (row, column), texts in authored_texts.items():
			lines = max(_wrapped_lines(text, widths[column]) for text in texts)
			if lines > 1:
				height = 15 * lines + 3
				dimension = sheet.row_dimensions[row]
				dimension.height = max(dimension.height or 15, height)
		# A brief written interpretation needs eight lines even when no answer is
		# supplied. Share that minimum across the authored rows so a tall response
		# area keeps its existing footprint rather than expanding every row.
		for region in source.role_regions:
			if region.role != "interpretation":
				continue
			reference = region.reference
			minimum_height = (15 * 8 + 3) / (reference.end_row - reference.start_row + 1)
			for row in range(reference.start_row, reference.end_row + 1):
				dimension = sheet.row_dimensions[row]
				dimension.height = max(dimension.height or 15, minimum_height)
	book.save(output_path)
	book.close()
	_preserve_numbers(document, pathlib.Path(output_path))


#============================================
def _content_bounds(source: tutorial_lib.model.Worksheet, start_row: int, start_column: int,
	end_row: int, end_column: int, include_entries: bool) -> tuple[int, int, int, int] | None:
	"""Find the real visual footprint of supplied content and student targets."""
	cells = []
	for row in range(start_row, end_row + 1):
		for column in range(start_column, end_column + 1):
			seed = source.cells.get((row, column))
			if ((seed is not None and seed.scalar.value not in (None, ""))
				or include_entries and (row, column) in source.entry_cells):
				cells.append((row, column))
	if not cells:
		return None
	rows = [item[0] for item in cells]
	columns = [item[1] for item in cells]
	return min(rows), min(columns), max(rows), max(columns)


#============================================
def _future_headers(document: tutorial_lib.model.Document, source: tutorial_lib.model.Worksheet,
		region_bounds: list[tuple[tuple[int, int, int, int], str, int]]) -> list[tuple[int, int, int]]:
	"""Find authored learner label targets that form table headers."""
	labels = set()
	for action in document.actions:
		if action.kind != "label" or action.target is None or action.target.sheet != source.name:
			continue
		labels.add((action.target.start_row, action.target.start_column))
	header_ranges = []
	eligible = set()
	for bounds, role, _ in region_bounds:
		if role == "source":
			continue
		row, start_column, _, end_column = bounds
		for candidate_row in range(bounds[0], bounds[2] + 1):
			for candidate_column in range(bounds[1], bounds[3] + 1):
				eligible.add((candidate_row, candidate_column))
		columns = [column for column in range(start_column, end_column + 1)
			if (row, column) in labels]
		while columns:
			first = columns.pop(0)
			last = first
			while columns and columns[0] == last + 1:
				last = columns.pop(0)
			header_ranges.append((row, first, last))
	# A contiguous run of two or more requested labels is a conventional header
	# even when it starts below the first row of its declared calculation area.
	# This recognizes authored ANOVA subtable headings without a special case.
	for row in sorted({item[0] for item in labels}):
		columns = sorted(column for label_row, column in labels
			if label_row == row and (label_row, column) in eligible)
		while columns:
			first = columns.pop(0)
			last = first
			while columns and columns[0] == last + 1:
				last = columns.pop(0)
			if last > first:
				header_ranges.append((row, first, last))
	return header_ranges


#============================================
def _unique_header_ranges(ranges: list[tuple[int, int, int]]) -> list[tuple[int, int, int]]:
	"""Keep overlapping source and planned headers from being styled twice."""
	result = []
	for item in ranges:
		if item not in result:
			result.append(item)
	return result


#============================================
def _source_stripe(source: tutorial_lib.model.Worksheet, row: int, column: int) -> bool:
	"""Tint alternate supplied-data rows without pretending empty padding is data."""
	for region in source.regions:
		if (region.start_row < row <= region.end_row
			and region.start_column <= column <= region.end_column):
			return (row - region.start_row) % 2 == 0
	return False


#============================================
def _entry_style(document: tutorial_lib.model.Document, sheet_name: str, row: int, column: int,
	role: str) -> tuple[str, str]:
	"""Make formula entry and label construction visible without revealing answers."""
	for action in document.actions:
		if (action.target is None or action.target.sheet != sheet_name
			or (action.target.start_row, action.target.start_column) != (row, column)):
			continue
		if action.kind == "label":
			return LABEL_ENTRY_STYLE
		if action.scalar is not None and action.scalar.kind == "formula":
			return FORMULA_ENTRY_STYLE
	background, accent, _ = ROLE_STYLES[role]
	return background, accent


#============================================
def _apply_region_outline(sheet: openpyxl.worksheet.worksheet.Worksheet, start_row: int,
		start_column: int, end_row: int, end_column: int, color: str) -> None:
	"""Give an authored block a clear boundary while keeping interior spacing light."""
	line = openpyxl.styles.Side(style="medium", color=color)
	for row in range(start_row, end_row + 1):
		for column in range(start_column, end_column + 1):
			cell = sheet.cell(row, column)
			cell.border = openpyxl.styles.Border(
				left=line if column == start_column else cell.border.left,
				right=line if column == end_column else cell.border.right,
				top=line if row == start_row else cell.border.top,
				bottom=line if row == end_row else cell.border.bottom)


#============================================
def _action_texts(action: tutorial_lib.model.Action) -> list[str]:
	"""Estimate text needs from literal entries and possible formula text operands."""
	texts = []
	if action.scalar.kind == "text":
		texts.append(action.scalar.value)
	elif action.scalar.kind == "formula":
		# Quoted operands include IF decisions; they are estimates, not calculated results.
		for token in openpyxl.formula.tokenizer.Tokenizer(action.scalar.value).items:
			if token.type == "OPERAND" and token.subtype == "TEXT":
				texts.append(token.value[1:-1].replace('""', '"'))
	return texts


#============================================
def _wrapped_lines(text: str, width: float) -> int:
	"""Reserve padded lines using the same character units as column sizing."""
	lines = sum(max(1, len(textwrap.wrap(line, width=max(1, int(width) - 2))))
		for line in text.split("\n"))
	return lines


#============================================
def _preserve_numbers(document: tutorial_lib.model.Document, path: pathlib.Path) -> None:
	"""Avoid openpyxl's 16-digit formatting changing a finite Python numeric value."""
	parts = _sheet_parts(path)
	with zipfile.ZipFile(path) as archive:
		members = [(item, archive.read(item.filename)) for item in archive.infolist()]
	updated = {}
	for sheet, part in zip(document.worksheets, parts, strict=True):
		root = defusedxml.ElementTree.fromstring(dict(
			(item.filename, content) for item, content in members)[part])
		for cell in root.iter(f"{MAIN_NS}c"):
			row, column = openpyxl.utils.cell.coordinate_to_tuple(cell.attrib["r"])
			seed = sheet.cells.get((row, column))
			if seed is not None and seed.scalar.kind == "number":
				cell.find(f"{MAIN_NS}v").text = repr(seed.scalar.value)
		updated[part] = defusedxml.ElementTree.tostring(root, encoding="utf-8")
	with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
		for item, content in members:
			archive.writestr(item, updated.get(item.filename, content))


#============================================
def verify_workbook(document: tutorial_lib.model.Document, path: pathlib.Path | str) -> None:
	"""Raise on changed worksheet order, seed values, types, or unexpected content."""
	names, observed = _logical_cells(pathlib.Path(path))
	if names != [sheet.name for sheet in document.worksheets]:
		raise ValueError("Workbook verification failed: worksheet order or names changed")
	blank = tutorial_lib.model.Scalar("blank", None)
	for sheet, cells in zip(document.worksheets, observed, strict=True):
		for address, seed in sheet.cells.items():
			expected = seed.scalar
			if expected.kind == "formula":
				expected = tutorial_lib.model.Scalar("formula",
					tutorial_lib.formula_io.logical_formula(expected.value))
			if cells.get(address, blank) != expected:
				raise ValueError(f"Workbook verification failed: {sheet.name}!{address} changed")
		for address, scalar in cells.items():
			if address not in sheet.cells and scalar.kind != "blank":
				raise ValueError(f"Workbook verification failed: unexpected {sheet.name}!{address}")


#============================================
def _inventory(path: pathlib.Path) -> None:
	"""Block content whose meaning cannot survive the supported Markdown format."""
	issues = []
	blocked_tags = {
		"mergeCells", "legacyDrawing", "legacyDrawingHF", "oleObjects",
		"controls", "tableParts", "pivotTableParts", "dataValidations", "conditionalFormatting",
		"hyperlinks", "extLst", "sheetProtection", "autoFilter", "scenarios",
		"definedName",
	}
	with zipfile.ZipFile(path) as archive:
		for name in archive.namelist():
			if any(piece in name for piece in (
				"vbaProject", "xl/charts/", "xl/pivot", "xl/externalLinks/", "xl/comments",
				"xl/threadedComments/", "xl/embeddings/", "xl/media/")):
				issues.append(f"unsupported package content: {name}")
			if name.startswith("xl/drawings/") and name.endswith(".xml"):
				# Excel emits empty drawing containers; these contain no lost content.
				if len(defusedxml.ElementTree.fromstring(archive.read(name))):
					issues.append(f"unsupported drawing content: {name}")
			if name == "xl/workbook.xml" or name.startswith("xl/worksheets/sheet"):
				if not name.endswith(".xml"):
					continue
				root = defusedxml.ElementTree.fromstring(archive.read(name))
				for item in root.iter():
					tag = item.tag.rsplit("}", 1)[-1]
					if tag in blocked_tags:
						issues.append(f"{name}: unsupported {tag}")
					if tag == "workbookProtection" and item.attrib:
						issues.append(f"{name}: workbook protection")
					if (tag == "calcPr"
						and item.attrib.get("iterate", "0").lower() in ("1", "true")):
						issues.append(
							"workbook enables iterative calculation; formula behavior cannot be preserved safely")
					if tag == "sheet" and item.attrib.get("state", "visible") != "visible":
						issues.append(f"hidden worksheet: {item.attrib['name']}")
					if tag in ("row", "col") and item.attrib.get("hidden") in ("1", "true"):
						issues.append(f"{name}: hidden {tag} content")
					if tag == "f" and item.attrib.get("t", "normal") not in ("normal", "shared"):
						issues.append(f"{name}: unsupported formula type {item.attrib['t']}")
					if tag == "c" and item.attrib.get("t") in ("e", "d"):
						issues.append(f"{name}!{item.attrib['r']}: error or date cell")
			if name == "xl/sharedStrings.xml" or name.startswith("xl/worksheets/sheet"):
				if name.endswith(".xml"):
					root = defusedxml.ElementTree.fromstring(archive.read(name))
					if any(True for item in root.iter(f"{MAIN_NS}r")):
						issues.append(f"{name}: rich text")
	if issues:
		message = "XLSX import blocked; supply a simpler workbook or extend the format:\n"
		message += "\n".join(sorted(set(issues)))
		raise ValueError(message)
	# Date styles may encode numeric cells as dates, even without the ISO date type.
	book = openpyxl.load_workbook(path, data_only=False)
	if book.epoch == openpyxl.utils.datetime.CALENDAR_MAC_1904:
		issues.append("workbook uses the 1904 date epoch; numeric serials cannot be preserved safely")
	storage = _xml_cells(path)
	for sheet, part in zip(book.worksheets, _sheet_parts(path), strict=True):
		for address in storage[part]:
			if sheet[address].is_date:
				issues.append(f"{sheet.title}!{address}: date/time cell")
		if sheet.title.startswith("'") or sheet.title.endswith("'"):
			issues.append(f"worksheet name cannot be represented: {sheet.title!r}")
		if (sheet.title != sheet.title.strip() or "\n" in sheet.title or "\r" in sheet.title
			or re.search(r"\s#+\s*$", sheet.title)):
			issues.append(f"worksheet name needs normalization before import: {sheet.title!r}")
	book.close()
	if issues:
		raise ValueError("XLSX import blocked; simplify these cells:\n" + "\n".join(issues))


#============================================
def _markdown_scalar(scalar: tutorial_lib.model.Scalar) -> str:
	"""Encode a scalar inside a pipe table without changing its semantic type."""
	if scalar.kind == "blank":
		return ""
	if scalar.kind == "boolean":
		return "TRUE" if scalar.value else "FALSE"
	if scalar.kind == "number":
		return str(scalar.value)
	value = json.dumps(scalar.value, ensure_ascii=True) if scalar.kind == "text" else scalar.value
	# Code spans preserve literal Markdown, entities and backslashes. GFM pipes still escape.
	if scalar.kind == "formula" and ("\n" in value or "\r" in value):
		raise ValueError("Multiline formulas require a simpler single-line source before import")
	runs = [len(piece) for piece in re.findall(r"`+", value)]
	fence = "`" * (max(runs, default=0) + 1)
	value = value.replace("|", "\\|")
	result = f"{fence} {value} {fence}"
	return result


#============================================
def _row_runs(cells: dict) -> list[tuple[int, int, int, int]]:
	"""Bound occupied consecutive row runs without allocating empty worksheet extents."""
	occupied = {key: value for key, value in cells.items() if value.kind != "blank"}
	if not occupied:
		return [(1, 1, 1, 1)]
	row_bounds = {}
	for row, column in occupied:
		low, high = row_bounds.get(row, (column, column))
		row_bounds[row] = (min(low, column), max(high, column))
	regions = []
	for row in sorted(row_bounds):
		low, high = row_bounds[row]
		if regions and regions[-1][2] + 1 == row:
			start, left, end, right = regions[-1]
			regions[-1] = (start, min(left, low), row, max(right, high))
		else:
			regions.append((row, low, row, high))
	if len(regions) == 1:
		start_row, start_column, end_row, end_column = regions[0]
		# Keep small near-A1 tables intuitive while retaining all original coordinates.
		if start_row <= 10 and start_column <= 10:
			start_row, start_column = 1, 1
		regions[0] = (start_row, start_column, end_row, end_column)
	area = sum((end_row - row + 1) * (end_column - column + 1)
		for row, column, end_row, end_column in regions)
	if area > MAX_IMPORT_CELLS:
		raise ValueError(
			f"XLSX sparse extents require {area} table cells; limit is {MAX_IMPORT_CELLS}. "
			"Split widely separated columns into simpler worksheets before importing.")
	return regions


#============================================
def import_workbook(path: pathlib.Path | str) -> str:
	"""Create a draft single-source document only after a loss-prevention inventory."""
	path = pathlib.Path(path)
	_inventory(path)
	names, sheets = _logical_cells(path)
	title = path.stem.replace("_", " ").replace("\n", " ")
	lines = [f"# {title}", "", "> Import summary: Draft migration; instructions require authoring.",
		"> Ordinary fonts, fills, widths, and number formats are discarded.",
		"> Formulas are preserved; calculation was not performed and cached results are not imported.",
		"", "## Workbook", ""]
	blank = tutorial_lib.model.Scalar("blank", None)
	for name, cells in zip(names, sheets, strict=True):
		lines.extend([f"### Sheet: {name}", ""])
		for index, (row, column, end_row, end_column) in enumerate(_row_runs(cells)):
			if index > 0 or (row, column) != (1, 1):
				address = f"{openpyxl.utils.get_column_letter(column)}{row}"
				lines.extend([f"#### At: {address}", ""])
			for source_row in range(row, end_row + 1):
				values = [_markdown_scalar(cells.get((source_row, col), blank))
					for col in range(column, end_column + 1)]
				lines.append("| " + " | ".join(values) + " |")
				if source_row == row:
					lines.append("| " + " | ".join("---" for value in values) + " |")
			lines.append("")
	lines.extend(["## Tutorial", "", "Write the tutorial instructions here.", ""])
	result = "\n".join(lines)
	return result
